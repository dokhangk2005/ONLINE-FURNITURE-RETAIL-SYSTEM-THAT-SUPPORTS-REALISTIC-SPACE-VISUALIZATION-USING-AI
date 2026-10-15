from __future__ import annotations

from pathlib import Path
from datetime import date

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\KLTN\khoaluan\artifacts\Tai_lieu_phan_bien_web_noi_that.docx")
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
PALE_BLUE = "F4F8FC"
WHITE = "FFFFFF"
GRAY = "666666"
CONTENT_DXA = 9360


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_keep_with_next(paragraph, value=True) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    tag = p_pr.find(qn("w:keepNext"))
    if value and tag is None:
        p_pr.append(OxmlElement("w:keepNext"))
    elif not value and tag is not None:
        p_pr.remove(tag)


def set_repeatable_row_no_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fld = OxmlElement("w:fldSimple")
    fld.set(qn("w:instr"), "PAGE")
    paragraph._p.append(fld)


def set_table_borders(table, color="B8C5D1", size="4") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:color"), color)


def configure_styles(doc: Document) -> None:
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for style_name, size, color, before, after in (
        ("Title", 28, DARK_BLUE, 0, 14),
        ("Subtitle", 13, GRAY, 0, 8),
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = style_name != "Subtitle"
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("List Bullet", "List Number"):
        style = doc.styles[style_name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def configure_section(section, first=False) -> None:
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)
    section.different_first_page_header_footer = first


def add_header_footer(section, title: str) -> None:
    for header in (section.header, section.even_page_header):
        p = header.paragraphs[0]
        p.text = title
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.runs[0].font.size = Pt(8.5)
        p.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    for footer in (section.footer, section.even_page_footer):
        add_page_number(footer.paragraphs[0])


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    set_keep_with_next(p)
    return p


def add_p(doc: Document, text: str = "", bold_lead: str | None = None):
    p = doc.add_paragraph()
    if bold_lead and text.startswith(bold_lead):
        r1 = p.add_run(bold_lead)
        r1.bold = True
        p.add_run(text[len(bold_lead):])
    else:
        p.add_run(text)
    return p


def _restart_numbering(doc: Document, paragraphs) -> None:
    numbering = doc.part.numbering_part.element
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    abstract_id = None
    for abstract in numbering.findall("w:abstractNum", ns):
        for style in abstract.findall(".//w:pStyle", ns):
            if style.get(qn("w:val")) == "ListNumber":
                abstract_id = abstract.get(qn("w:abstractNumId"))
                break
        if abstract_id is not None:
            break
    if abstract_id is None:
        abstract_id = "0"
    current_ids = [int(n.get(qn("w:numId"))) for n in numbering.findall("w:num", ns)]
    num_id = max(current_ids, default=0) + 1
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_ref = OxmlElement("w:abstractNumId")
    abstract_ref.set(qn("w:val"), str(abstract_id))
    num.append(abstract_ref)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:startOverride")
    start.set(qn("w:val"), "1")
    override.append(start)
    num.append(override)
    numbering.append(num)
    for paragraph in paragraphs:
        p_pr = paragraph._p.get_or_add_pPr()
        num_pr = p_pr.get_or_add_numPr()
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), "0")
        num_id_node = OxmlElement("w:numId")
        num_id_node.set(qn("w:val"), str(num_id))
        num_pr.append(ilvl)
        num_pr.append(num_id_node)


def add_bullets(doc: Document, items: list[str], numbered=False) -> None:
    style = "List Number" if numbered else "List Bullet"
    paragraphs = []
    for item in items:
        paragraphs.append(doc.add_paragraph(item, style=style))
    if numbered:
        _restart_numbering(doc, paragraphs)


def add_note(doc: Document, title: str, text: str, fill=PALE_BLUE) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_width(cell, CONTENT_DXA)
    set_cell_margins(cell, top=140, bottom=140, start=180, end=180)
    set_cell_shading(cell, fill)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = cell.add_paragraph(text)
    p2.paragraph_format.space_after = Pt(0)
    set_table_borders(table, color=LIGHT_BLUE, size="6")


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int], font_size=9.2):
    assert sum(widths) == CONTENT_DXA
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(CONTENT_DXA))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    set_repeatable_row_no_split(hdr)
    for idx, text in enumerate(headers):
        cell = hdr.cells[idx]
        set_cell_width(cell, widths[idx])
        set_cell_margins(cell)
        set_cell_shading(cell, LIGHT_BLUE)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(text)
        run.bold = True
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor.from_string(DARK_BLUE)

    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        set_repeatable_row_no_split(table.rows[-1])
        for idx, text in enumerate(row):
            cell = cells[idx]
            set_cell_width(cell, widths[idx])
            set_cell_margins(cell)
            if ridx % 2 == 1:
                set_cell_shading(cell, "FAFCFE")
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            run = p.add_run(str(text))
            run.font.size = Pt(font_size)
            if idx == 0 and len(headers) >= 3:
                run.bold = True
        for cell in cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return table


def add_function(doc: Document, title: str, purpose: str, steps: list[str], notes: str | None = None) -> None:
    add_heading(doc, title, 3)
    add_p(doc, purpose)
    add_bullets(doc, steps, numbered=True)
    if notes:
        p = doc.add_paragraph()
        r = p.add_run("Điểm kỹ thuật: ")
        r.bold = True
        p.add_run(notes)


def add_page_break(doc: Document) -> None:
    doc.add_page_break()


def api_rows():
    groups: list[tuple[str, list[tuple[str, str, str, str]]]] = [
        ("Xác thực và tài khoản", [
            ("POST", "/api/auth/quick-register", "Công khai", "Gửi liên kết đăng ký nhanh qua email"),
            ("GET", "/api/auth/registration-link", "Công khai", "Kiểm tra token thiết lập mật khẩu"),
            ("POST", "/api/auth/complete-registration", "Công khai", "Đặt mật khẩu và hoàn tất tài khoản"),
            ("POST", "/api/auth/login", "Công khai", "Đăng nhập và nhận JWT"),
            ("POST", "/api/auth/forgot-password", "Công khai", "Gửi OTP khôi phục mật khẩu"),
            ("POST", "/api/auth/verify-otp", "Công khai", "Xác minh OTP"),
            ("POST", "/api/auth/reset-password", "Công khai", "Đặt lại mật khẩu"),
            ("GET", "/api/users/admin/all", "Admin", "Danh sách tài khoản"),
            ("PUT", "/api/users/admin/update-role/{userId}", "Admin", "Cập nhật vai trò"),
            ("PUT", "/api/users/admin/toggle-status/{userId}", "Admin", "Khóa/mở tài khoản"),
            ("PUT", "/api/users/admin/reset-password/{userId}", "Admin", "Đặt lại mật khẩu tài khoản"),
            ("PUT", "/api/users/update-profile/{userId}", "Đăng nhập", "Cập nhật hồ sơ cá nhân"),
        ]),
        ("Danh mục, sản phẩm và ảnh", [
            ("GET", "/api/categories", "Công khai", "Lấy cây danh mục"),
            ("POST", "/api/categories", "Admin", "Tạo danh mục"),
            ("PUT", "/api/categories/{id}", "Admin", "Sửa danh mục"),
            ("DELETE", "/api/categories/{id}", "Admin", "Xóa mềm danh mục"),
            ("GET", "/api/products?includeDeleted=false", "Công khai", "Danh sách sản phẩm; Admin có thể gồm bản ghi xóa"),
            ("GET", "/api/products/{id}", "Công khai", "Chi tiết sản phẩm/biến thể/ảnh"),
            ("POST", "/api/products", "Admin", "Tạo sản phẩm và biến thể"),
            ("PUT", "/api/products/{id}", "Admin", "Sửa sản phẩm và biến thể"),
            ("DELETE", "/api/products/{id}", "Admin", "Xóa mềm sản phẩm"),
            ("POST", "/api/upload/image", "Admin, Shipper", "Tải ảnh; dùng cho sản phẩm/bằng chứng"),
            ("GET", "/uploads/products/by-id/{imageId}", "Công khai", "Trả ảnh sản phẩm theo ID"),
        ]),
        ("Địa chỉ và bản đồ", [
            ("GET", "/api/addresses", "Đăng nhập", "Danh sách địa chỉ đã lưu"),
            ("POST", "/api/addresses", "Đăng nhập", "Thêm địa chỉ"),
            ("PUT", "/api/addresses/{id}", "Đăng nhập", "Sửa địa chỉ"),
            ("DELETE", "/api/addresses/{id}", "Đăng nhập", "Xóa địa chỉ"),
            ("GET", "/api/location/suggest", "Đăng nhập", "Proxy gợi ý địa chỉ Mapbox"),
            ("GET", "/api/location/retrieve/{mapboxId}", "Đăng nhập", "Lấy chi tiết địa điểm/tọa độ"),
        ]),
        ("Voucher", [
            ("GET", "/api/vouchers/admin", "Admin", "Danh sách voucher quản trị"),
            ("GET", "/api/vouchers/active", "Công khai", "Danh sách voucher đang hiệu lực"),
            ("POST", "/api/vouchers/validate", "Công khai", "Kiểm tra và tính giảm giá"),
            ("POST", "/api/vouchers/admin", "Admin", "Tạo voucher"),
            ("PUT", "/api/vouchers/admin/{id}", "Admin", "Sửa voucher"),
            ("DELETE", "/api/vouchers/admin/{id}", "Admin", "Xóa voucher"),
        ]),
        ("Đơn hàng và thanh toán", [
            ("POST", "/api/orders", "Đăng nhập", "Tạo đơn; trừ tồn kho; tạo thanh toán"),
            ("GET", "/api/orders/vnpay-return", "VNPay/Công khai", "Return URL, kiểm tra chữ ký"),
            ("GET", "/api/orders/vnpay-ipn", "VNPay/Công khai", "IPN cập nhật thanh toán"),
            ("GET", "/api/orders/{orderId}/payment-status", "Đăng nhập", "Tra trạng thái thanh toán"),
            ("GET", "/api/orders/my-orders/{userId}", "Đăng nhập", "Lịch sử đơn của người dùng"),
            ("GET", "/api/orders/admin/all", "Admin", "Danh sách toàn bộ đơn"),
            ("PUT", "/api/orders/admin/{orderId}/status", "Admin", "Đổi trạng thái trong giới hạn nghiệp vụ"),
            ("GET", "/api/orders/shippers", "Admin", "Danh sách shipper để điều phối"),
            ("POST", "/api/orders/calculate-shipping", "Đăng nhập", "Tính phí theo khoảng cách"),
            ("POST", "/api/orders/admin/{orderId}/assign-shipper", "Admin", "Phân công shipper"),
            ("GET", "/api/orders/shipper/{shipperId}", "Shipper, Admin", "Danh sách đơn của shipper"),
            ("PUT", "/api/orders/shipper/{orderId}/complete", "Shipper", "Hoàn tất giao sau các điều kiện bắt buộc"),
        ]),
        ("Vận chuyển, thu tiền và đối soát", [
            ("GET", "/api/delivery/admin-overview", "Admin", "Tổng quan vận chuyển/shipper/đơn"),
            ("GET", "/api/delivery/shipper/availability", "Shipper", "Xem trạng thái sẵn sàng"),
            ("PUT", "/api/delivery/shipper/availability", "Shipper", "Bật/tắt sẵn sàng nhận đơn"),
            ("GET", "/api/delivery/shipper/profile", "Shipper", "Xem hồ sơ giao hàng"),
            ("PUT", "/api/delivery/shipper/profile", "Shipper", "Cập nhật xe, khu vực, năng lực"),
            ("PUT", "/api/delivery/shipper/shift", "Shipper", "Đổi ca: sẵn sàng/đang giao/nghỉ/offline"),
            ("POST", "/api/delivery/shipper/{orderId}/accept", "Shipper", "Nhận đơn được phân công"),
            ("POST", "/api/delivery/shipper/{orderId}/reject", "Shipper", "Từ chối và trả đơn về Admin"),
            ("PUT", "/api/delivery/shipper/{orderId}/status", "Shipper", "Tiến trạng thái giao hàng"),
            ("POST", "/api/delivery/shipper/{orderId}/otp/request", "Shipper", "Gửi OTP giao hàng cho khách"),
            ("POST", "/api/delivery/shipper/{orderId}/otp/verify", "Shipper", "Xác minh OTP giao hàng"),
            ("POST", "/api/delivery/shipper/{orderId}/payment/collect", "Shipper", "Xác nhận thu tiền và bằng chứng"),
            ("GET", "/api/delivery/shipper/reconciliation", "Shipper", "Tổng kết khoản đã thu"),
            ("POST", "/api/delivery/shipper/reconciliation/submit", "Shipper", "Gửi bàn giao tiền"),
            ("GET", "/api/delivery/admin/reconciliations", "Admin", "Danh sách đối soát"),
            ("POST", "/api/delivery/admin/reconciliation/{paymentId}/confirm", "Admin", "Xác nhận đã nhận tiền"),
            ("POST", "/api/delivery/admin/reconciliation/{paymentId}/reject", "Admin", "Trả lại yêu cầu đối soát"),
            ("GET", "/api/delivery/shipper/performance", "Shipper", "Hiệu suất ngày/tuần/tháng"),
            ("GET", "/api/delivery/admin/shipper-performance", "Admin", "So sánh hiệu suất shipper"),
            ("PUT", "/api/delivery/shipper/{orderId}/report-issue", "Shipper", "Báo sự cố có ghi chú/bằng chứng"),
            ("PUT", "/api/delivery/shipper/{orderId}/cancel", "Shipper", "Hủy giao theo điều kiện thanh toán"),
            ("PUT", "/api/delivery/admin/{orderId}/reschedule", "Admin", "Đặt lại lịch giao"),
        ]),
        ("Đánh giá", [
            ("POST", "/api/reviews", "Đăng nhập", "Đánh giá sản phẩm đã mua/giao thành công"),
            ("GET", "/api/reviews/product/{productId}", "Công khai", "Danh sách đánh giá được duyệt"),
            ("GET", "/api/reviews/product/{productId}/stats", "Công khai", "Điểm trung bình và phân bố sao"),
            ("GET", "/api/reviews/admin/pending", "Admin", "Danh sách chờ duyệt"),
            ("PUT", "/api/reviews/admin/{reviewId}/approve", "Admin", "Duyệt đánh giá"),
            ("DELETE", "/api/reviews/admin/{reviewId}", "Admin", "Từ chối/xóa đánh giá"),
            ("GET", "/api/reviews/admin/products", "Admin", "Tổng hợp đánh giá theo sản phẩm"),
            ("GET", "/api/reviews/admin/products/{productId}", "Admin", "Chi tiết đánh giá sản phẩm"),
            ("PUT", "/api/reviews/admin/{reviewId}/reply", "Admin", "Phản hồi và gửi email thông báo"),
        ]),
        ("Hỏi đáp cũ và hỗ trợ thời gian thực", [
            ("POST", "/api/questions/create", "Đăng nhập", "Tạo câu hỏi kiểu cũ"),
            ("GET", "/api/questions/user/{userId}", "Đăng nhập", "Câu hỏi của người dùng"),
            ("GET", "/api/questions/admin", "Admin", "Danh sách câu hỏi kiểu cũ"),
            ("PUT", "/api/questions/admin/{id}/answer", "Admin", "Trả lời và gửi email"),
            ("GET", "/api/support-conversations/mine", "Đăng nhập", "Hộp thư hội thoại của khách"),
            ("GET", "/api/support-conversations/admin", "Admin", "Hộp thư quản trị, lọc trạng thái"),
            ("GET", "/api/support-conversations/{id}", "Đăng nhập", "Nội dung hội thoại"),
            ("POST", "/api/support-conversations", "Đăng nhập", "Tạo hội thoại; có rate limit"),
            ("POST", "/api/support-conversations/{id}/messages", "Đăng nhập", "Gửi tin; chống gửi trùng"),
            ("PUT", "/api/support-conversations/{id}/read", "Đăng nhập", "Đánh dấu đã đọc"),
            ("PUT", "/api/support-conversations/{id}/status", "Admin", "Đóng/mở hội thoại"),
            ("PUT", "/api/support-conversations/{id}/assign", "Admin", "Admin tự nhận xử lý"),
        ]),
        ("Chatbot", [
            ("POST", "/api/chatbot/ask", "Đăng nhập", "Gửi câu hỏi; Gemini chính, OpenAI dự phòng"),
        ]),
        ("Gợi ý sản phẩm", [
            ("GET", "/api/recommendations/admin/dashboard", "Admin", "Thống kê, cấu hình và luật gợi ý"),
            ("PUT", "/api/recommendations/admin/settings", "Admin", "Bật/tắt thuật toán và ngưỡng"),
            ("POST", "/api/recommendations/admin/rebuild", "Admin", "Chạy lại Apriori"),
            ("DELETE", "/api/recommendations/admin/rules", "Admin", "Xóa các luật đã sinh"),
            ("POST", "/api/recommendations/run-apriori", "Admin", "API tương thích cũ để chạy Apriori"),
            ("GET", "/api/recommendations/frequently-bought/{productId}", "Công khai", "Sản phẩm thường mua cùng"),
            ("GET", "/api/recommendations/similar/{productId}", "Công khai", "Sản phẩm tương tự content-based"),
        ]),
        ("Doanh thu", [
            ("GET", "/api/revenue/admin/report", "Admin", "Báo cáo theo from/to/groupBy=day|month|year"),
        ]),
        ("Phối cảnh không gian", [
            ("GET", "/api/room-designs/harmonization/status", "Công khai", "Kiểm tra nhà cung cấp AI tạo ảnh"),
            ("GET", "/api/room-designs/{id}/harmonized/{cacheKey}", "Token thiết kế", "Trả ảnh AI đã sinh/cache"),
            ("POST", "/api/room-designs/{id}/generate-preview", "Token thiết kế", "Sinh ảnh từ phòng + tối đa 5 ảnh sản phẩm"),
            ("POST", "/api/room-designs", "Công khai/nhận JWT nếu có", "Tạo thiết kế và lưu ảnh phòng"),
            ("GET", "/api/room-designs/mine", "Đăng nhập", "Thiết kế của người dùng"),
            ("GET", "/api/room-designs/{id}", "Token thiết kế", "Lấy metadata/scene"),
            ("GET", "/api/room-designs/{id}/room", "Token thiết kế", "Trả ảnh phòng gốc"),
            ("GET", "/api/room-designs/{id}/preview", "Token thiết kế", "Trả ảnh preview thủ công"),
            ("GET", "/api/room-designs/{id}/analysis", "Token thiết kế", "Kết quả phân tích sàn/tường/chiều sâu"),
            ("GET", "/api/room-designs/{id}/analysis/{kind}", "Token thiết kế", "Ảnh depth/floor/wall/overlay/occlusion"),
            ("POST", "/api/room-designs/{id}/analyze", "Token thiết kế", "Gọi Python AI và cache kết quả"),
            ("PUT", "/api/room-designs/{id}", "Token thiết kế", "Lưu scene JSON và tiêu đề"),
            ("POST", "/api/room-designs/{id}/preview", "Token thiết kế", "Lưu ảnh preview"),
            ("DELETE", "/api/room-designs/{id}", "Token thiết kế", "Xóa thiết kế và file liên quan"),
            ("POST", "/api/visualization/products/{imageId}/remove-background", "Công khai", "Gọi rembg xóa nền sản phẩm"),
            ("GET", "/api/visualization/products/{imageId}/transparent", "Công khai", "Trả ảnh PNG nền trong"),
            ("GET", "/api/visualization/health", "Công khai", "Kiểm tra Python AI service"),
        ]),
    ]
    return groups


def build_document() -> None:
    doc = Document()
    doc.settings.odd_and_even_pages_header_footer = True
    configure_styles(doc)
    configure_section(doc.sections[0], first=True)
    add_header_footer(doc.sections[0], "HỒ SƠ PHẢN BIỆN • MY FURNITURE STORE")

    # Cover: editorial_cover adapted to a clean academic defense handout.
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(88)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run("TÀI LIỆU PHẢN BIỆN ĐỒ ÁN")
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor.from_string(BLUE)

    title = doc.add_paragraph(style="Title")
    title.add_run("WEB THƯƠNG MẠI ĐIỆN TỬ\nNỘI THẤT THÔNG MINH")
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.add_run("My Furniture Store • Flutter Web + ASP.NET Core + AI")

    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    cell = table.cell(0, 0)
    set_cell_width(cell, CONTENT_DXA)
    set_cell_shading(cell, LIGHT_BLUE)
    set_cell_margins(cell, top=240, bottom=240, start=260, end=260)
    set_table_borders(table, color=LIGHT_BLUE, size="0")
    for i, line in enumerate([
        "Mục đích: trình bày chức năng, kiến trúc, công nghệ, cách sử dụng và hệ thống API.",
        "Phạm vi đọc source: Flutter Web tại D:\\KLTN\\khoaluan; ASP.NET Core và Python AI tại D:\\KLTN\\furniture_BE\\furniture_BE.",
        "Ảnh chụp source: 16/08/2026. Tài liệu không chứa API key, mật khẩu hoặc chuỗi kết nối.",
    ]):
        p = cell.paragraphs[0] if i == 0 else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.add_run(line)

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(210)
    r = p.add_run("BẢN DÙNG CHO BUỔI PHẢN BIỆN")
    r.bold = True
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor.from_string(DARK_BLUE)
    p2 = doc.add_paragraph("Ngày lập: 16/08/2026")
    p2.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
    add_page_break(doc)

    add_heading(doc, "Mục lục nội dung", 1)
    toc = [
        "1. Tóm tắt đồ án và điểm nổi bật",
        "2. Kiến trúc tổng thể và công nghệ",
        "3. Chức năng dành cho khách hàng",
        "4. Chức năng dành cho Admin",
        "5. Chức năng dành cho Shipper",
        "6. Các luồng nghiệp vụ quan trọng",
        "7. AI, thuật toán gợi ý và phối cảnh",
        "8. Cơ sở dữ liệu và bảo mật",
        "9. Hướng dẫn chạy và trình diễn",
        "10. Danh mục API nội bộ và tích hợp ngoài",
        "11. Kiểm thử, giới hạn và câu hỏi phản biện",
    ]
    add_bullets(doc, toc, numbered=False)
    add_note(doc, "Cách dùng tài liệu", "Nếu chỉ có 10–15 phút chuẩn bị, đọc mục 1, sơ đồ ở mục 2, kịch bản demo ở mục 9 và câu hỏi phản biện ở mục 11. Phần API là phụ lục để tra cứu khi hội đồng hỏi sâu.")

    add_heading(doc, "1. Tóm tắt đồ án và điểm nổi bật", 1)
    add_p(doc, "My Furniture Store là web thương mại điện tử nội thất có ba vai trò: Khách hàng, Admin và Shipper. Hệ thống không chỉ bán hàng mà còn xử lý hậu cần, thu tiền/đối soát, hỗ trợ thời gian thực và phối cảnh không gian bằng AI.")
    add_table(doc, ["Khối", "Giá trị chính", "Điểm nổi bật khi bảo vệ"], [
        ["Mua sắm", "Sản phẩm nhiều biến thể, bộ lọc, giỏ hàng, voucher, thanh toán", "Giá/tồn kho được tính lại ở server; hỗ trợ VNPay sandbox"],
        ["Cá nhân hóa", "Apriori và content-based", "Admin điều chỉnh ngưỡng, bật/tắt và rebuild luật"],
        ["Phối cảnh", "Thủ công 2D + sinh ảnh AI", "Phân tích sàn/tường/depth; ghép tối đa 5 sản phẩm"],
        ["Vận chuyển", "Điều phối, OTP, bằng chứng, thu tiền, đối soát", "Luồng trạng thái có điều kiện và SignalR realtime"],
        ["Chăm sóc", "Chatbot + hội thoại khách–Admin", "SignalR, typing/read receipt, email thông báo"],
        ["Quản trị", "Doanh thu, đơn, tài khoản, đánh giá, voucher", "Báo cáo ngày/tháng/năm theo sản phẩm và danh mục"],
        ["Ngôn ngữ", "Tiếng Việt và tiếng Anh", "Lưu lựa chọn bằng SharedPreferences"],
    ], [1500, 3600, 4260], font_size=9.5)

    add_heading(doc, "Điểm khác biệt của đồ án", 2)
    add_bullets(doc, [
        "Tích hợp chuỗi mua hàng đầy đủ từ xem sản phẩm đến giao hàng và đối soát tiền.",
        "Tính năng phối cảnh có hai hướng: người dùng tự điều chỉnh và AI sinh ảnh tham khảo từ ảnh phòng + ảnh sản phẩm.",
        "Gợi ý sản phẩm có thể giải thích được: Apriori dựa trên giỏ hàng đã giao; content-based dựa trên thuộc tính sản phẩm.",
        "Đánh giá chỉ dành cho người đã nhận sản phẩm; nội dung được kiểm duyệt tự động và Admin có thể phản hồi qua email.",
        "Shipper có dashboard riêng, ca làm việc, trạng thái sẵn sàng, bằng chứng giao, thu tiền và chỉ số hiệu suất.",
    ])
    add_note(doc, "Thông điệp trình bày 30 giây", "Đây là hệ thống bán nội thất full-stack, tập trung vào ba bài toán khó: trải nghiệm xem thử sản phẩm trong phòng, cá nhân hóa gợi ý và quản lý giao hàng/thu tiền có kiểm soát.")

    add_heading(doc, "2. Kiến trúc tổng thể và công nghệ", 1)
    add_heading(doc, "2.1 Sơ đồ kiến trúc", 2)
    add_table(doc, ["Lớp giao diện", "Lớp nghiệp vụ", "Dữ liệu & AI", "Dịch vụ ngoài"], [[
        "Flutter Web\nMaterial UI\nVI/EN\nSharedPreferences",
        "ASP.NET Core Web API\nJWT/RBAC\nSignalR\nRate limiting",
        "SQL Server + EF Core\nPython FastAPI\nONNX/rembg\nKho file ảnh",
        "VNPay\nMapbox\nSMTP Gmail\nOpenAI/Gemini",
    ]], [2200, 2450, 2450, 2260], font_size=9.4)
    add_p(doc, "Luồng chính: trình duyệt gọi REST API để đọc/ghi dữ liệu; JWT xác định danh tính và vai trò; SignalR đẩy sự kiện chat/giao hàng; backend gọi SQL Server, Python AI hoặc dịch vụ ngoài tùy nghiệp vụ.")

    add_heading(doc, "2.2 Công nghệ sử dụng", 2)
    add_table(doc, ["Tầng", "Công nghệ", "Vai trò trong hệ thống"], [
        ["Frontend", "Flutter / Dart SDK ^3.11.5", "Một codebase web responsive, Material UI, điều hướng và state cục bộ"],
        ["Giao tiếp HTTP", "http 1.6.0, http_parser", "REST, JSON, multipart upload"],
        ["Trình duyệt", "web 1.1.1, url_launcher", "Download ảnh, mở VNPay/Google Maps/liên kết"],
        ["Lưu cục bộ", "shared_preferences 2.5.3", "Phiên đăng nhập, ngôn ngữ, giỏ hàng theo người dùng"],
        ["Realtime", "signalr_netcore 1.4.4", "Chat hỗ trợ và cập nhật vận chuyển"],
        ["Backend", "ASP.NET Core .NET 10", "REST API, middleware, DI, auth, file serving"],
        ["ORM/CSDL", "EF Core 10.0.9 + SQL Server", "Migrations, quan hệ, truy vấn và transaction"],
        ["Bảo mật", "JWT Bearer 10.0.9 + BCrypt.Net-Next 4.2.0", "Xác thực, phân quyền, băm mật khẩu"],
        ["AI nội bộ", "Python FastAPI + Pillow + rembg + ONNX Runtime", "Xóa nền, phân đoạn sàn/tường, chiều sâu"],
        ["AI ngoài", "OpenAI Images/Moderation; Gemini/OpenAI chat", "Sinh ảnh phối cảnh, lọc tục tĩu, chatbot"],
        ["Thanh toán", "VNPay sandbox", "Tạo URL thanh toán, callback/IPN, xác thực chữ ký"],
        ["Bản đồ", "Mapbox Search", "Gợi ý địa chỉ và tọa độ để tính khoảng cách"],
        ["Email", "SMTP Gmail", "Đăng ký, OTP, xác nhận đơn, phản hồi đánh giá/hỏi đáp"],
    ], [1450, 3020, 4890], font_size=9.2)

    add_heading(doc, "2.3 Tổ chức source", 2)
    add_table(doc, ["Thư mục", "Nội dung"], [
        ["lib/page", "Màn hình khách hàng, Admin, Shipper, checkout và visualizer"],
        ["lib/model", "Model JSON cho sản phẩm, đơn, doanh thu, chat, phối cảnh"],
        ["lib/services", "ApiService, AuthService, CartService, SignalR, locale"],
        ["lib/widgets", "Navbar/footer, chatbot, địa chỉ, lớp chỉnh mask và widget dùng chung"],
        ["Controllers", "19 controller, 116 action REST"],
        ["Services", "Email, recommendation, moderation, VNPay, SignalR notifier, AI image"],
        ["Models + Data", "Entity SQL Server và AppDbContext"],
        ["ai_service", "FastAPI, mô hình ONNX và xử lý ảnh"],
    ], [2400, 6960], font_size=9.5)

    add_page_break(doc)
    add_heading(doc, "3. Chức năng dành cho khách hàng", 1)
    add_function(doc, "3.1 Đăng ký nhanh và đăng nhập", "Giảm số bước đăng ký nhưng không gửi mật khẩu thô qua email.", [
        "Nhập email và bấm Đăng ký nhanh.",
        "Hệ thống gửi liên kết thiết lập mật khẩu dùng một lần, hiệu lực 30 phút.",
        "Mở liên kết, đặt mật khẩu rồi đăng nhập bằng email/tên tài khoản.",
        "Nếu quên mật khẩu, dùng luồng OTP 6 số qua email để đặt lại.",
    ], "Token đăng ký được băm SHA-256; mật khẩu lưu bằng BCrypt; đăng nhập trả JWT có hạn dùng cấu hình.")
    add_function(doc, "3.2 Trang chủ, tìm kiếm và duyệt sản phẩm", "Giúp khách khám phá nội thất theo danh mục và nhu cầu.", [
        "Xem banner, danh mục và sản phẩm bán chạy.",
        "Dùng tìm kiếm trên thanh điều hướng hoặc mở trang Sản phẩm.",
        "Lọc theo danh mục, giá, màu, kích thước; sắp xếp danh sách.",
        "Mở chi tiết để xem ảnh gốc, biến thể màu/chất liệu/kích thước, giá và tồn kho.",
    ], "Ảnh hiển thị bán hàng là ảnh gốc; ảnh xóa nền chỉ phục vụ phối cảnh.")
    add_function(doc, "3.3 Gợi ý sản phẩm", "Tăng khả năng khám phá và bán chéo.", [
        "Ở chi tiết sản phẩm, xem mục Thường mua cùng dựa trên Apriori.",
        "Xem Sản phẩm tương tự dựa trên danh mục, chất liệu, màu và độ gần giá.",
        "Chọn biến thể phù hợp trước khi thêm vào giỏ.",
    ], "Apriori chỉ học từ đơn đủ điều kiện đã giao/thanh toán thành công; content-based có thể hoạt động ngay cả khi dữ liệu đơn còn ít.")
    add_function(doc, "3.4 Giỏ hàng và voucher", "Quản lý các biến thể muốn mua và khuyến mãi.", [
        "Đăng nhập để mở giỏ hàng; giỏ được tách theo tài khoản.",
        "Tăng/giảm số lượng trong giới hạn tồn kho hoặc xóa sản phẩm.",
        "Chọn voucher trong Kho ưu đãi hoặc nhập mã.",
        "Hệ thống kiểm tra thời gian hiệu lực, đơn tối thiểu, giới hạn dùng và mức giảm tối đa.",
    ])
    add_function(doc, "3.5 Đặt hàng và thanh toán", "Thu thập thông tin nhận hàng và tạo đơn an toàn.", [
        "Chọn địa chỉ đã lưu hoặc gõ địa chỉ để nhận gợi ý Mapbox.",
        "Chọn thời gian giao, có/không lắp ráp, voucher và phương thức thanh toán.",
        "Backend tính lại giá, giảm giá, phí vận chuyển và trừ tồn kho trong transaction.",
        "Với VNPay: chuyển sang sandbox, nhận return/IPN và xác minh checksum.",
        "Với POS/COD: thanh toán chỉ thành công sau khi shipper xác nhận đã thu; chuyển khoản ở trạng thái chờ xác nhận.",
    ], "Phí đang mã hóa trong source: 50.000đ cho 5 km đầu + 8.000đ/km vượt 5 km (làm tròn lên); lắp ráp 100.000đ.")
    add_function(doc, "3.6 Tài khoản, địa chỉ và lịch sử đơn", "Khách tự quản lý thông tin sau bán hàng.", [
        "Cập nhật họ tên, số điện thoại, ngày sinh và thông tin cá nhân.",
        "Thêm/sửa/xóa địa chỉ; chọn nhãn Nhà/Công ty/Khác và địa chỉ mặc định.",
        "Xem lịch sử đơn, trạng thái thanh toán và vận chuyển.",
        "Chỉ đánh giá sản phẩm từ đơn đã giao; mỗi khách chỉ một đánh giá cho mỗi sản phẩm.",
    ])
    add_function(doc, "3.7 Hỗ trợ và chatbot", "Cung cấp kênh tự phục vụ và trao đổi trực tiếp.", [
        "Hỏi chatbot về sản phẩm/chính sách; backend ưu tiên Gemini và có nhánh OpenAI dự phòng.",
        "Mở hội thoại hỗ trợ với Admin; nhận tin nhắn theo thời gian thực.",
        "Theo dõi trạng thái đã đọc, đang nhập và hội thoại mở/đóng.",
        "Nhận email khi Admin phản hồi đánh giá hoặc hỏi đáp.",
    ], "SignalR hub hỗ trợ reconnect; API gửi tin có rate limit và ClientMessageId chống gửi trùng.")
    add_function(doc, "3.8 Đổi ngôn ngữ", "Chuyển giao diện giữa tiếng Việt và tiếng Anh.", [
        "Bấm nút ngôn ngữ ở giao diện khách, đăng nhập, Admin hoặc Shipper.",
        "Chọn VI/EN; ứng dụng cập nhật locale ngay.",
        "Lựa chọn được lưu để dùng lại ở lần mở sau.",
    ])

    add_heading(doc, "4. Chức năng dành cho Admin", 1)
    admin_rows = [
        ["Dashboard", "Số liệu tổng quan và biểu đồ; biểu đồ cột tập trung tỷ lệ đơn giao thành công."],
        ["Danh mục", "Tạo, sửa, xóa mềm; hỗ trợ quan hệ danh mục cha–con."],
        ["Sản phẩm", "CRUD sản phẩm, mã, mô tả, danh mục; nhiều biến thể và nhiều ảnh; chọn thumbnail."],
        ["Đơn hàng", "Tìm/lọc/xem chi tiết; chuyển trạng thái hợp lệ; không vượt quyền shipper khi đang giao."],
        ["Tài khoản", "Tìm/lọc, phân vai trò Admin/Customer/Shipper, khóa/mở và reset mật khẩu."],
        ["Voucher", "Tạo/sửa/xóa, phần trăm hoặc số tiền cố định, thời gian, hạn mức và đơn tối thiểu."],
        ["Đánh giá", "Xem theo sản phẩm, duyệt/từ chối, phản hồi; moderation AI chặn nội dung tục tĩu."],
        ["Vận chuyển", "Xem workload và trạng thái shipper; điều phối, đặt lại lịch, bằng chứng và sự cố."],
        ["Đối soát", "Duyệt hoặc trả lại khoản shipper đã thu; lưu lịch sử thay đổi."],
        ["Q&A realtime", "Hộp thư hội thoại, tìm/lọc, tự nhận xử lý, trả lời, đóng/mở."],
        ["Gợi ý", "Bật/tắt Apriori/content-based, đặt support/confidence/max results, rebuild và xóa luật."],
        ["Doanh thu", "Lọc khoảng ngày; nhóm ngày/tháng/năm; tổng hợp sản phẩm, danh mục, phí và giảm giá."],
    ]
    add_table(doc, ["Phân hệ", "Khả năng quản trị"], admin_rows, [2000, 7360], font_size=9.4)

    add_heading(doc, "Quy tắc doanh thu", 2)
    add_bullets(doc, [
        "Chỉ công nhận đơn có OrderStatus = Delivered; thời điểm ghi nhận là UpdatedAt, nếu thiếu dùng OrderDate.",
        "Tách doanh thu hàng hóa, phí vận chuyển, phí lắp ráp và ước tính giảm giá.",
        "Tính số đơn, số sản phẩm, khách hàng duy nhất, giá trị đơn trung bình và số đơn dùng voucher.",
        "So sánh kỳ hiện tại với kỳ liền trước có cùng độ dài.",
        "Bảng sản phẩm có số lượng, số đơn, doanh thu, giá bán trung bình và tỷ trọng; bảng danh mục có số sản phẩm và tỷ trọng.",
        "Giới hạn truy vấn: theo ngày tối đa 370 ngày, theo tháng 10 năm, theo năm 40 năm.",
    ])

    add_heading(doc, "5. Chức năng dành cho Shipper", 1)
    add_table(doc, ["Nhóm", "Chức năng", "Điều kiện/kiểm soát"], [
        ["A. Hồ sơ", "Xe, biển số, khu vực, giao hàng cồng kềnh, hỗ trợ lắp ráp, số đơn tối đa", "Admin dùng để phân công phù hợp"],
        ["B. Ca làm", "Sẵn sàng, đang giao, tạm nghỉ, ngoại tuyến", "Không mặc định hoạt động 24/24"],
        ["C. Nhận đơn", "Xem đơn được giao; chấp nhận hoặc từ chối có lý do", "Từ chối trả đơn về Admin"],
        ["D. Hành trình", "Gọi khách, mở Google Maps, cập nhật trạng thái", "Trạng thái tiến theo thứ tự nghiệp vụ"],
        ["E. Giao hàng", "Đến nơi, gửi/xác minh OTP, ảnh bằng chứng, tên người nhận", "Đủ điều kiện mới được hoàn tất"],
        ["F. Sự cố/hủy", "Báo sự cố, ghi chú, ảnh; hủy đơn theo điều kiện", "Đơn đã thu/đã trả tiền cần kiểm soát Admin"],
        ["G. Thu tiền", "Hiển thị số phải thu, xác nhận đã thu, ảnh bằng chứng, gửi đối soát", "Không tự đổi paid khi chưa thu"],
        ["H. Hiệu suất", "Đang giao, hoàn thành, thất bại, km, đúng hẹn, hẹn lại, thời gian TB", "Thống kê ngày/tuần/tháng"],
    ], [1350, 4430, 3580], font_size=9.2)
    add_heading(doc, "Luồng thao tác giao một đơn", 2)
    add_bullets(doc, [
        "Bật Sẵn sàng và vào ca.",
        "Chấp nhận đơn Admin đã phân công.",
        "Cập nhật lần lượt: đã nhận hàng → đang giao → đã đến nơi.",
        "Nếu POS/COD, xác nhận số tiền thực thu và bằng chứng trước khi hoàn tất.",
        "Gửi OTP cho khách, nhập OTP, tải ảnh giao hàng và ghi nhận người nhận.",
        "Hoàn tất; hệ thống đẩy realtime và cập nhật báo cáo hiệu suất.",
        "Cuối ca, gửi các khoản đã thu để Admin xác nhận bàn giao.",
    ], numbered=True)

    add_page_break(doc)
    add_heading(doc, "6. Các luồng nghiệp vụ quan trọng", 1)
    add_heading(doc, "6.1 Tạo đơn và đảm bảo nhất quán", 2)
    add_table(doc, ["Bước", "Xử lý phía server"], [
        ["1", "Xác thực người dùng và request."],
        ["2", "Đọc lại biến thể, giá và tồn kho từ SQL Server; không tin giá frontend gửi lên."],
        ["3", "Xác thực voucher và tính giảm giá."],
        ["4", "Tính phí vận chuyển/lắp ráp và tổng tiền."],
        ["5", "Tạo Order, OrderDetail, Payment và trừ tồn kho trong transaction mức Serializable."],
        ["6", "Nếu VNPay, tạo URL ký HMAC; nếu lỗi thanh toán, hủy đơn và hoàn tồn/voucher."],
        ["7", "Gửi email xác nhận và trả mã đơn dạng DH-yyyyMMdd-xxxxxx."],
    ], [900, 8460], font_size=9.4)
    add_heading(doc, "6.2 Trạng thái giao hàng", 2)
    add_p(doc, "Luồng trạng thái được bảo vệ ở backend, không chỉ ẩn nút ở giao diện. Admin không thể tự đánh dấu Delivered khi shipper đã nhận nhưng chưa hoàn thành các bước bắt buộc.")
    add_table(doc, ["Giai đoạn", "Dữ liệu minh chứng"], [
        ["Phân công", "ShipperId, thời điểm phân công, trạng thái chấp nhận/từ chối"],
        ["Nhận hàng", "AcceptedAt/PickedUpAt và trạng thái vận chuyển"],
        ["Đến nơi", "ArrivedAt"],
        ["Xác nhận khách", "OTP đã xác minh, thời điểm xác minh"],
        ["Hoàn tất", "Ảnh bằng chứng, người nhận, DeliveredAt"],
        ["Sự cố/hủy", "Lý do, ghi chú, ảnh, người thao tác, thời điểm"],
    ], [2200, 7160], font_size=9.4)
    add_heading(doc, "6.3 Đánh giá có xác thực mua hàng", 2)
    add_bullets(doc, [
        "Backend kiểm tra khách có OrderDetail của sản phẩm thuộc đơn Delivered.",
        "Ràng buộc unique (ProductId, UserId) ngăn đánh giá trùng.",
        "Danh sách từ khóa cục bộ chặn nhanh; OpenAI Moderation đánh giá bổ sung.",
        "Nội dung an toàn có thể tự duyệt; nội dung không chắc chắn chuyển sang chờ Admin; nội dung vi phạm bị chặn.",
        "Admin phản hồi sẽ gửi email cho khách.",
    ])
    add_heading(doc, "6.4 Hỗ trợ thời gian thực", 2)
    add_table(doc, ["Kênh", "REST", "SignalR"], [
        ["Khách", "Tạo hội thoại, tải lịch sử, gửi tin, đánh dấu đọc", "Nhận tin, trạng thái hội thoại, typing và read receipt"],
        ["Admin", "Tải inbox, tự nhận, đổi trạng thái, gửi phản hồi", "Cập nhật inbox và tin nhắn tức thời"],
        ["Vận chuyển", "Cập nhật trạng thái/OTP/thu tiền/sự cố", "Đẩy DeliveryUpdated để các màn hình tải lại đúng lúc"],
    ], [1400, 3900, 4060], font_size=9.4)

    add_heading(doc, "7. AI, thuật toán gợi ý và phối cảnh", 1)
    add_heading(doc, "7.1 Apriori", 2)
    add_p(doc, "Mục tiêu là tìm các cặp sản phẩm thường xuất hiện trong cùng đơn hàng đủ điều kiện. Mỗi đơn là một basket; hệ thống đếm tần suất cặp, tính support và confidence, rồi lưu luật có hướng Product → RecommendedProduct.")
    add_table(doc, ["Khái niệm", "Ý nghĩa"], [
        ["Support", "Tỷ lệ giỏ hàng chứa đồng thời cặp sản phẩm; lọc các cặp đủ phổ biến."],
        ["Confidence", "Xác suất mua B khi đã mua A; dùng để xếp hạng luật có hướng."],
        ["MinSupport/MinConfidence", "Ngưỡng do Admin cấu hình."],
        ["MaxResults", "Số gợi ý tối đa trả về."],
    ], [2600, 6760], font_size=9.5)
    add_heading(doc, "7.2 Content-based", 2)
    add_p(doc, "Tính điểm tương tự bằng thuộc tính sản phẩm: cùng danh mục +4, cùng chất liệu +2, cùng màu +1; sau đó ưu tiên sản phẩm có giá gần sản phẩm hiện tại. Ưu điểm là không cần nhiều lịch sử giao dịch.")

    add_heading(doc, "7.3 Phối cảnh thủ công", 2)
    add_bullets(doc, [
        "Tải ảnh phòng JPG/JPEG/PNG/WEBP tối đa 10 MB.",
        "Backend lưu thiết kế; Python AI trả bản đồ chiều sâu, mask sàn/tường, overlay, foreground/occlusion và thông tin ánh sáng.",
        "Chọn sản phẩm/biến thể; ảnh được xóa nền bằng rembg và thêm vào canvas.",
        "Kéo thả, pinch/slider đổi tỷ lệ, xoay phẳng, xoay phối cảnh (yaw), lật ngang.",
        "Chọn bám sàn, tường trái, tường phải hoặc tự do; vị trí được snap vào mặt phẳng phân tích.",
        "Tỷ lệ có thể thay đổi theo depth; chỉnh độ sáng và bóng để hòa cảnh.",
        "Lưu scene JSON và xuất ảnh PNG độ phân giải 2x.",
    ])
    add_note(doc, "Phạm vi chính xác hiện tại", "Source có model và widget chỉnh mask che khuất, backend cũng sinh occlusion mask; tuy nhiên trang RoomVisualizerPage hiện chưa gắn widget chỉnh mask vào luồng giao diện chính và đang lưu danh sách occlusionEdits rỗng. Khi bảo vệ nên nói đây là nền tảng kỹ thuật đã có, còn thao tác sửa mask trực tiếp cần hoàn thiện tích hợp UI.", fill="FFF4E5")

    add_heading(doc, "7.4 Tạo ảnh phối cảnh bằng AI", 2)
    add_bullets(doc, [
        "Khách tải ảnh phòng và chuyển sang chế độ AI.",
        "Chọn từ 1 đến 5 sản phẩm/biến thể; danh sách hiển thị dạng ngăn xếp.",
        "Có thể nhập yêu cầu vị trí tối đa 500 ký tự.",
        "Backend gửi ảnh phòng + các ảnh sản phẩm đến nhà cung cấp đã cấu hình (OpenAI là cấu hình hiện tại; có provider Gemini thay thế).",
        "Kết quả được cache theo nội dung phòng, sản phẩm, model và prompt để giảm chi phí/thời gian.",
        "Khách xem, tạo lại và tải ảnh tham khảo.",
    ])
    add_p(doc, "Thông điệp cần nói rõ: ảnh AI là hình minh họa tham khảo, không bảo đảm kích thước vật lý chính xác như AR/3D đo đạc thực tế.")

    add_heading(doc, "7.5 Python AI service", 2)
    add_table(doc, ["Thành phần", "Công nghệ / vai trò"], [
        ["API", "FastAPI + Uvicorn, endpoint /health, /remove-background, /analyze-room"],
        ["Xóa nền", "rembg[cpu], mô hình u2netp, kết quả PNG trong suốt"],
        ["Chiều sâu", "Depth Anything V2 ViT-S ONNX"],
        ["Phân đoạn", "SegFormer-B0 ADE20K ONNX để xác định sàn/tường/foreground"],
        ["Xử lý ảnh", "Pillow + NumPy"],
        ["Suy luận", "ONNX Runtime CPU; không bắt buộc GPU"],
        ["Tối ưu", "Lưu cache file, khóa theo thiết kế để tránh gọi trùng"],
    ], [2500, 6860], font_size=9.4)

    add_page_break(doc)
    add_heading(doc, "8. Cơ sở dữ liệu và bảo mật", 1)
    add_heading(doc, "8.1 Các nhóm bảng chính", 2)
    add_table(doc, ["Nhóm", "Entity/bảng", "Quan hệ đáng chú ý"], [
        ["Người dùng", "Roles, Users, UserAddresses", "1 Role–n User; 1 User–n Address/Order"],
        ["Sản phẩm", "Categories, Products, ProductVariants, ProductImages", "Danh mục cây; sản phẩm nhiều biến thể/ảnh"],
        ["Bán hàng", "Orders, OrderDetails, Payments, Vouchers", "OrderDetail chốt UnitPrice; Order có Payment"],
        ["Vận chuyển", "Orders, Payments, PaymentReconciliationHistories", "Lưu shipper, mốc giao, bằng chứng, lịch sử đối soát"],
        ["Tương tác", "Reviews, SupportQuestions", "Review unique theo user + product"],
        ["Chat realtime", "SupportConversations, SupportMessages", "Tin nhắn có ClientMessageId unique chống trùng"],
        ["Gợi ý", "RecommendationSettings, ProductRecommendations", "Luật có hướng, Restrict delete tránh cascade path"],
        ["Phối cảnh", "RoomDesigns + kho file", "Scene JSON, token truy cập, hạn dùng và cache ảnh"],
    ], [1500, 3420, 4440], font_size=9.1)
    add_heading(doc, "8.2 Cơ chế bảo mật", 2)
    add_bullets(doc, [
        "JWT Bearer xác thực API; [Authorize(Roles=...)] phân quyền Admin và Shipper.",
        "BCrypt băm mật khẩu; token đăng ký lưu dạng hash và có thời hạn/sử dụng một lần.",
        "VNPay callback cho phép anonymous nhưng bắt buộc kiểm tra vnp_SecureHash, mã giao dịch và số tiền.",
        "Giao dịch tạo đơn dùng isolation Serializable để tránh bán vượt tồn kho.",
        "Soft delete cho User, Category, Product và ProductVariant.",
        "CORS chỉ cho các origin local đã cấu hình khi phát triển; static file có cache policy và nosniff.",
        "Rate limiting cho kênh hỗ trợ; SignalR xác thực từ access token.",
        "Room design chưa đăng nhập dùng access token ngẫu nhiên 32 byte; tài khoản đăng nhập có thể xem mục mine.",
        "API key và mật khẩu phải ở User Secrets/biến môi trường, không commit vào appsettings.json.",
    ])
    add_note(doc, "Lưu ý an toàn", "Không trình chiếu file secrets.json trước hội đồng. Nếu một API key từng xuất hiện trong chat, ảnh hoặc lịch sử Git, nên thu hồi và tạo key mới trước khi demo.", fill="FFF4E5")

    add_heading(doc, "9. Hướng dẫn chạy và trình diễn", 1)
    add_heading(doc, "9.1 Điều kiện", 2)
    add_bullets(doc, [
        ".NET SDK 10, Flutter SDK tương thích Dart 3.11.5, SQL Server và Python 3.",
        "User Secrets của backend có ConnectionStrings:DefaultConnection, Jwt:Key và các khóa tích hợp cần dùng.",
        "Cài dependency Python từ ai_service/requirements.txt và đặt model ONNX đúng cấu trúc service.",
        "Mặc định frontend gọi backend http://127.0.0.1:5024; backend gọi Python http://127.0.0.1:8001.",
    ])
    add_heading(doc, "9.2 Trình tự chạy", 2)
    add_table(doc, ["Thứ tự", "Lệnh / thao tác", "Mục đích"], [
        ["1", "Khởi động SQL Server", "Đảm bảo database truy cập được"],
        ["2", "cd D:\\KLTN\\furniture_BE\\furniture_BE\\ai_service; .\\run.ps1", "AI xóa nền/phân tích phòng ở cổng 8001"],
        ["3", "cd D:\\KLTN\\furniture_BE\\furniture_BE; dotnet run --project furniture_BE.csproj", "REST API + SignalR ở cổng cấu hình"],
        ["4", "cd D:\\KLTN\\khoaluan; flutter run -d web-server --web-port 8080", "Flutter Web"],
        ["5", "Mở đúng URL mà Flutter in ra", "Tránh dùng link/port của lần chạy cũ"],
    ], [950, 5540, 2870], font_size=8.8)
    add_p(doc, "Script ai_service/run.ps1 chạy app.py bằng Uvicorn; nếu service đã chạy thì kiểm tra GET /health trước khi demo phối cảnh.")

    add_heading(doc, "9.3 Kịch bản demo 8–10 phút", 2)
    demo = [
        "Khách: đổi ngôn ngữ, tìm/lọc sản phẩm, chọn biến thể và xem gợi ý.",
        "Khách: thêm giỏ, chọn địa chỉ gợi ý, áp voucher và tạo đơn (ưu tiên POS khi không muốn phụ thuộc sandbox).",
        "Admin: xem đơn, điều phối một shipper đang Sẵn sàng.",
        "Shipper: nhận đơn, cập nhật trạng thái, OTP/bằng chứng, thu tiền và hoàn tất.",
        "Admin: xem doanh thu Delivered và xác nhận đối soát.",
        "Khách: từ lịch sử đơn gửi đánh giá; Admin duyệt/phản hồi.",
        "Khách–Admin: gửi tin realtime ở Trung tâm hỗ trợ.",
        "Phối cảnh: tải ảnh phòng, thử thủ công; sau đó chọn 1–2 sản phẩm và sinh ảnh AI nếu quota ổn định.",
    ]
    add_bullets(doc, demo, numbered=True)
    add_heading(doc, "9.4 Phương án dự phòng khi demo", 2)
    add_bullets(doc, [
        "Nếu OpenAI/Gemini bận hoặc hết quota: trình diễn chế độ thủ công và ảnh AI đã cache.",
        "Nếu Mapbox lỗi: dùng địa chỉ đã lưu có tọa độ hoặc nhập đầy đủ rồi tiếp tục.",
        "Nếu VNPay sandbox lỗi: dùng POS/COD và giải thích luồng callback ký số.",
        "Nếu SignalR mất kết nối: tải lại lịch sử qua REST; giao diện có trạng thái kết nối/retry.",
        "Chuẩn bị sẵn một tài khoản cho mỗi role và một đơn ở từng mốc trạng thái.",
    ])

    add_heading(doc, "10. Danh mục API nội bộ và tích hợp ngoài", 1)
    add_p(doc, "Backend hiện có 19 controller và 116 action REST. Bảng dưới liệt kê endpoint theo source. {id} là tham số đường dẫn; quyền “Đăng nhập” yêu cầu JWT, còn “Token thiết kế” dùng accessToken riêng của RoomDesign.")
    for group_name, rows in api_rows():
        add_heading(doc, group_name, 2)
        add_table(doc, ["Method", "Endpoint", "Quyền", "Mục đích"], [list(r) for r in rows], [820, 3760, 1500, 3280], font_size=8.3)

    add_heading(doc, "10.1 SignalR hubs", 2)
    add_table(doc, ["Hub", "Khách sử dụng", "Sự kiện chính"], [
        ["/hubs/support", "Customer, Admin", "ConversationUpdated, MessageReceived, MessagesRead, TypingChanged"],
        ["/hubs/delivery", "Admin, Shipper, màn hình chung", "DeliveryUpdated và cập nhật điều phối/trạng thái"],
    ], [2200, 2600, 4560], font_size=9.2)

    add_heading(doc, "10.2 API Python nội bộ", 2)
    add_table(doc, ["Method", "Endpoint", "Input", "Output"], [
        ["GET", "http://127.0.0.1:8001/health", "—", "Trạng thái service/model"],
        ["POST", "/remove-background", "multipart image", "PNG nền trong"],
        ["POST", "/analyze-room", "multipart room image", "Metadata + depth/floor/wall/overlay/occlusion images"],
    ], [900, 2880, 2250, 3330], font_size=9.2)

    add_heading(doc, "10.3 Tích hợp API/dịch vụ ngoài", 2)
    add_table(doc, ["Dịch vụ", "Endpoint/cơ chế", "Dùng cho", "Khóa cấu hình"], [
        ["VNPay Sandbox", "paymentv2/vpcpay.html + return/IPN", "Thanh toán VND và checksum", "VnPay:TmnCode, HashSecret, Url, ReturnUrl"],
        ["Mapbox", "Search suggest/retrieve", "Gợi ý địa chỉ và tọa độ", "Mapbox:AccessToken"],
        ["OpenAI Images", "POST /v1/images/edits", "Sinh ảnh phối cảnh", "OpenAI:ApiKey, ImageModel/Quality/Size"],
        ["OpenAI Moderation", "POST /v1/moderations", "Lọc đánh giá tục tĩu", "OpenAI:ApiKey, ModerationModel"],
        ["OpenAI Chat", "Chat API", "Fallback chatbot", "OpenAI:ApiKey, ChatModel"],
        ["Gemini", "models/{model}:generateContent", "Chatbot; provider ảnh thay thế", "Gemini:ApiKey, ChatModel, ImageModel"],
        ["SMTP Gmail", "SMTP TLS cổng 587", "Đăng ký, OTP, đơn, phản hồi", "SmtpSettings:*"],
    ], [1500, 2780, 2340, 2740], font_size=8.7)
    add_note(doc, "Nguyên tắc cấu hình", "Tài liệu chỉ nêu tên khóa. Giá trị bí mật phải ở .NET User Secrets hoặc biến môi trường; frontend không được chứa key OpenAI, Gemini, VNPay secret hay mật khẩu SMTP.")

    add_heading(doc, "11. Kiểm thử, giới hạn và câu hỏi phản biện", 1)
    add_heading(doc, "11.1 Trạng thái kiểm thử tại thời điểm lập tài liệu", 2)
    add_table(doc, ["Hạng mục", "Kết quả", "Ghi chú"], [
        ["flutter test", "40/40 test pass", "Chạy ngày 16/08/2026"],
        ["dotnet build --no-restore", "Build succeeded", "0 warning, 0 error"],
        ["Test frontend", "Auth/cart isolation, locale, review eligibility, phối cảnh, doanh thu, shipper, SignalR chat", "Có widget/model/service tests"],
        ["Backend", "Chưa thấy project test tự động riêng", "Nên bổ sung integration test cho thanh toán và state machine giao hàng"],
    ], [2200, 2200, 4960], font_size=9.2)

    add_heading(doc, "11.2 Giới hạn hiện tại và hướng phát triển", 2)
    add_bullets(doc, [
        "Phối cảnh thủ công là mô phỏng 2D có depth/perspective, chưa phải AR 3D có kích thước thật.",
        "Widget sửa occlusion mask chưa được nối lại vào màn visualizer hiện tại; cần hoàn thiện UI và lưu edit thay vì danh sách rỗng.",
        "AI sinh ảnh phụ thuộc quota, billing, độ trễ và chính sách của OpenAI/Gemini; cache chỉ giảm các request lặp.",
        "Hai phân hệ SupportQuestions cũ và SupportConversations realtime đang cùng tồn tại; nên migrate dứt điểm về hội thoại realtime.",
        "JWT lưu phía client bằng SharedPreferences phù hợp đồ án/demo; production nên cân nhắc cookie HttpOnly/SameSite và refresh token.",
        "Cần thêm backend integration test, kiểm thử tải SignalR, quan sát lỗi tập trung và quy trình backup file ảnh/database.",
        "Cấu hình localhost/CORS phải chuyển sang cấu hình môi trường khi deploy HTTPS thực tế.",
    ])

    add_heading(doc, "11.3 Câu hỏi phản biện thường gặp", 2)
    qa = [
        ["Tại sao dùng Flutter Web?", "Tái sử dụng component và logic Dart, responsive đa nền tảng; đổi lại bundle web lớn hơn web thuần và SEO cần giải pháp bổ sung."],
        ["Tại sao phải tính lại giá ở backend?", "Frontend có thể bị sửa; server phải là nguồn sự thật cho giá, voucher, phí và tồn kho."],
        ["Apriori khác content-based thế nào?", "Apriori học quan hệ mua cùng từ đơn; content-based so thuộc tính và hữu ích khi chưa đủ giao dịch."],
        ["Làm sao ngăn đánh giá giả?", "Chỉ user có đơn Delivered chứa sản phẩm mới được đánh giá; unique user–product ngăn gửi trùng."],
        ["VNPay có an toàn không?", "Secret chỉ ở backend; callback/IPN xác minh chữ ký, mã giao dịch và số tiền trước khi cập nhật."],
        ["Tại sao cần Python riêng?", "Hệ sinh thái model ảnh/ONNX/rembg mạnh; tách service giúp backend .NET giữ nghiệp vụ và scale AI riêng."],
        ["AI phối cảnh có chính xác không?", "Chỉ là hình tham khảo thẩm mỹ; độ chính xác kích thước cần 3D/AR, calibration và model sản phẩm."],
        ["Realtime mất mạng thì sao?", "REST vẫn là nguồn dữ liệu; client reconnect SignalR và có thể tải lại lịch sử/trạng thái."],
        ["Tại sao không cho Admin tự Delivered?", "Tách nhiệm vụ và yêu cầu OTP/bằng chứng/thu tiền giúp giảm gian lận và giữ audit trail."],
        ["Bảo vệ API key thế nào?", "User Secrets/biến môi trường, không commit, phân quyền và có kế hoạch rotate khi lộ."],
    ]
    add_table(doc, ["Câu hỏi", "Câu trả lời ngắn"], qa, [3300, 6060], font_size=9.0)

    add_heading(doc, "11.4 Checklist trước khi vào phòng phản biện", 2)
    add_bullets(doc, [
        "Kiểm tra SQL Server, backend, Python AI và Flutter đều chạy; mở sẵn tab /health hoặc log.",
        "Đăng nhập sẵn ba trình duyệt/profile cho Customer, Admin và Shipper.",
        "Có ít nhất một sản phẩm đủ ảnh/biến thể/tồn kho và một voucher còn hạn.",
        "Có đơn mẫu Pending, Assigned/In transit và Delivered để không mất thời gian thao tác dài.",
        "Kiểm tra VNPay sandbox, Mapbox, SMTP và quota OpenAI/Gemini; chuẩn bị phương án dự phòng.",
        "Ẩn file secrets.json, console chứa token và lịch sử chat có API key.",
        "Chuẩn bị một ảnh phòng và một ảnh AI đã cache để demo phối cảnh ổn định.",
    ])

    add_heading(doc, "Phụ lục: Nguồn đối chiếu chính", 1)
    add_table(doc, ["Nội dung", "File/source đã đối chiếu"], [
        ["Route và ứng dụng Flutter", "lib/main.dart, pubspec.yaml"],
        ["REST client", "lib/services/api_service.dart"],
        ["Đăng nhập/giỏ/ngôn ngữ", "lib/services/auth_service.dart, cart_service.dart, locale_service.dart"],
        ["Realtime", "lib/services/support_realtime_service.dart, delivery_realtime_service.dart"],
        ["Phối cảnh", "lib/page/visualizer/room_visualizer_page.dart; model/room_*; utils/placement_snap.dart"],
        ["Admin", "lib/page/admin/**"],
        ["Shipper", "lib/page/shipper/**"],
        ["Backend setup", "Program.cs, furniture_BE.csproj, appsettings.json"],
        ["API", "Controllers/*.cs"],
        ["CSDL", "Data/AppDbContext.cs, Models/*.cs, Migrations/*"],
        ["AI/gợi ý/email", "Services/*.cs, ai_service/*"],
    ], [2400, 6960], font_size=9.0)
    add_p(doc, "Kết luận: Source hiện thể hiện một hệ thống thương mại điện tử nội thất có phạm vi rộng và có tính tích hợp. Khi bảo vệ, nên nhấn mạnh tính nhất quán nghiệp vụ ở backend, phân quyền ba vai trò, khả năng giải thích của thuật toán gợi ý và giới hạn trung thực của ảnh phối cảnh AI.")

    # Metadata and document properties.
    doc.core_properties.title = "Tài liệu phản biện web nội thất thông minh"
    doc.core_properties.subject = "Chức năng, công nghệ, cách sử dụng và API"
    doc.core_properties.author = "Nhóm đồ án My Furniture Store"
    doc.core_properties.keywords = "Flutter, ASP.NET Core, SQL Server, AI, VNPay, SignalR, Apriori"
    doc.core_properties.comments = "Tổng hợp từ source ngày 16/08/2026; không chứa bí mật cấu hình."

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build_document()
