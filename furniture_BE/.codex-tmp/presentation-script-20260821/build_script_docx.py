from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(r"D:\KLTN\furniture_BE\KichBan_ThuyetTrinh_KhoaLuan_MyFurnitureStore.docx")
LOGO = Path(r"D:\KLTN\furniture_BE\.codex-tmp\title-slide-redesign-20260819\huflit-logo-white.png")

BLUE = RGBColor(61, 141, 255)
DARK = RGBColor(16, 24, 40)
MUTED = RGBColor(91, 100, 114)
LIGHT_BLUE = "EEF5FF"
LIGHT_GRAY = "F5F7FA"
RULE = "D7DEE8"


def set_run_font(run, name="Calibri", size=None, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def shade_paragraph(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def add_left_border(paragraph, color="3D8DFF", size="18", space="8"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), size)
    left.set(qn("w:space"), space)
    left.set(qn("w:color"), color)
    p_bdr.append(left)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run("Trang ")
    set_run_font(run, size=9, color=MUTED)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr_text)
    run._r.append(fld_char2)


def add_label_paragraph(doc, label, value, after=4, color=DARK):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=MUTED, bold=True)
    r2 = p.add_run(value)
    set_run_font(r2, size=10.5, color=color)
    return p


def add_body(doc, text, bold_lead=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.widow_control = True
    if bold_lead and text.startswith(bold_lead):
        lead = p.add_run(bold_lead)
        set_run_font(lead, size=11, color=DARK, bold=True)
        rest = p.add_run(text[len(bold_lead):])
        set_run_font(rest, size=11, color=DARK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=11, color=DARK)
    return p


def add_callout(doc, label, text, fill=LIGHT_BLUE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.13)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, fill)
    add_left_border(p)
    r1 = p.add_run(f"{label}: ")
    set_run_font(r1, size=10.5, color=BLUE, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=DARK, italic=True)
    return p


def add_action(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.line_spacing = 1.2
    shade_paragraph(p, LIGHT_GRAY)
    r1 = p.add_run("THAO TÁC: ")
    set_run_font(r1, size=10, color=MUTED, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=DARK)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    run = p.add_run(text)
    return p


def add_slide_section(doc, number, title, time_text, objective, paragraphs, transition=None, page_break=False, presenter_note=None):
    h = doc.add_paragraph(style="Heading 1")
    if page_break:
        h.paragraph_format.page_break_before = True
    run = h.add_run(f"SLIDE {number:02d} | {title}")
    set_run_font(run, size=16, color=BLUE, bold=True)
    add_label_paragraph(doc, "Thời lượng", time_text, after=2)
    if presenter_note:
        add_label_paragraph(doc, "Gợi ý phân công", presenter_note, after=5)
    add_callout(doc, "Mục tiêu", objective, fill=LIGHT_GRAY)
    h2 = doc.add_paragraph(style="Heading 2")
    r = h2.add_run("Lời trình bày gợi ý")
    set_run_font(r, size=13, color=DARK, bold=True)
    for text in paragraphs:
        add_body(doc, text)
    if transition:
        add_callout(doc, "Chuyển ý", transition)


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.82)
section.bottom_margin = Inches(0.78)
section.left_margin = Inches(0.9)
section.right_margin = Inches(0.9)
section.header_distance = Inches(0.38)
section.footer_distance = Inches(0.38)
section.different_first_page_header_footer = True

# Compact reference guide token map.
styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
normal.font.size = Pt(11)
normal.font.color.rgb = DARK
normal.paragraph_format.space_before = Pt(0)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for style_name, size, color, before, after in [
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, RGBColor(31, 77, 120), 10, 5),
]:
    st = styles[style_name]
    st.font.name = "Calibri"
    st._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    st._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    st.font.size = Pt(size)
    st.font.bold = True
    st.font.color.rgb = color
    st.paragraph_format.space_before = Pt(before)
    st.paragraph_format.space_after = Pt(after)
    # Avoid Word moving a chained heading block outside the printable area
    # when a slide section happens to begin close to a page boundary.
    st.paragraph_format.keep_with_next = False

# Running header and footer, hidden on cover page.
header = section.header
hp = header.paragraphs[0]
hp.paragraph_format.space_after = Pt(0)
hp.paragraph_format.tab_stops.add_tab_stop(Inches(6.7))
hr1 = hp.add_run("KỊCH BẢN THUYẾT TRÌNH KHÓA LUẬN")
set_run_font(hr1, size=8.5, color=MUTED, bold=True)
hr2 = hp.add_run("\tMY FURNITURE STORE")
set_run_font(hr2, size=8.5, color=MUTED, bold=True)

footer = section.footer
fp = footer.paragraphs[0]
add_page_number(fp)

# Cover page - editorial cover.
for _ in range(2):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(0)

if LOGO.exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(14)
    logo_shape = p.add_run().add_picture(str(LOGO), width=Inches(1.55))
    logo_shape._inline.docPr.set("title", "Logo HUFLIT")
    logo_shape._inline.docPr.set("descr", "Logo Trường Đại học Ngoại ngữ - Tin học Thành phố Hồ Chí Minh")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(12)
r = p.add_run("BẢO VỆ ĐỒ ÁN KHÓA LUẬN TỐT NGHIỆP")
set_run_font(r, size=11, color=BLUE, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(10)
r = p.add_run("KỊCH BẢN THUYẾT TRÌNH")
set_run_font(r, size=28, color=DARK, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(8)
r = p.add_run("Hệ thống bán lẻ nội thất trực tuyến\nhỗ trợ phối cảnh không gian thực bằng AI")
set_run_font(r, size=17, color=RGBColor(52, 64, 84), bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(20)
p.paragraph_format.space_after = Pt(5)
r = p.add_run("Thời lượng mục tiêu: khoảng 30 phút")
set_run_font(r, size=11, color=BLUE, bold=True)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_after = Pt(4)
r = p.add_run("Sinh viên: Đỗ Minh Vĩ Khang - Trần Ngọc Nhật Duy")
set_run_font(r, size=10.5, color=MUTED)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("GVHD: ThS. Lương Văn Minh | TP.HCM, 08/2026")
set_run_font(r, size=10.5, color=MUTED)

doc.add_page_break()

# Usage guide.
h = doc.add_paragraph(style="Heading 1")
r = h.add_run("CÁCH SỬ DỤNG KỊCH BẢN")
set_run_font(r, size=16, color=BLUE, bold=True)
add_label_paragraph(doc, "Cách luyện", "Đọc thành tiếng ít nhất hai lần; lần thứ ba chỉ nhìn các ý chính và tự diễn đạt.")
add_label_paragraph(doc, "Tốc độ", "Nói rõ ràng, khoảng 110-125 từ mỗi phút; dừng ngắn sau các kết luận quan trọng.")
add_label_paragraph(doc, "Khi trình bày", "Không quay lưng đọc slide. Chỉ tay vào sơ đồ hoặc màn hình đúng lúc nội dung được nhắc đến.")
add_label_paragraph(doc, "Phân công", "Có thể chia người trình bày tại slide 10 hoặc slide 11; người thao tác demo nên là người hiểu luồng nghiệp vụ nhất.")
add_callout(doc, "Phân bổ thời gian", "Slide 1-16 khoảng 21 phút; demo slide 17 khoảng 5 phút 30 giây; slide 18-20 khoảng 2 phút; phần còn lại dành cho chuyển cảnh và câu hỏi.")

h = doc.add_paragraph(style="Heading 2")
r = h.add_run("Nguyên tắc trả lời Hội đồng")
set_run_font(r, size=13, color=DARK, bold=True)
add_body(doc, "Mỗi câu trả lời nên đi theo ba bước: nêu quyết định thiết kế, chỉ ra bằng chứng trong báo cáo hoặc source code, sau đó nói rõ giới hạn hiện tại. Nếu chưa có số liệu, cần nói thẳng là nhóm chưa đánh giá thay vì suy đoán.")


add_slide_section(
    doc, 1, "Mở đầu và giới thiệu đề tài", "45 giây",
    "Tạo ấn tượng chuyên nghiệp, nêu đúng đề tài và giá trị cốt lõi trong hai câu đầu.",
    [
        "Kính thưa quý Thầy, Cô trong Hội đồng. Em là Đỗ Minh Vĩ Khang, cùng bạn Trần Ngọc Nhật Duy. Hôm nay nhóm em xin trình bày đề tài: Hệ thống bán lẻ nội thất trực tuyến hỗ trợ phối cảnh không gian thực bằng AI, dưới sự hướng dẫn của ThS. Lương Văn Minh.",
        "Đề tài của nhóm không chỉ xây dựng một website bán hàng. Mục tiêu chính là kết nối trải nghiệm chọn nội thất, tư vấn và xem thử sản phẩm trong phòng thật với quy trình quản lý đơn hàng, thanh toán và giao nhận ở phía sau.",
        "Trong phần trình bày, nhóm sẽ tập trung vào bài toán, kiến trúc giải pháp, cơ chế gợi ý, quy trình phối cảnh AI, kết quả triển khai và phần demo hệ thống.",
    ],
    "Trước hết, em xin trình bày ngắn mạch nội dung của buổi bảo vệ.",
    page_break=True,
)

add_slide_section(
    doc, 2, "Mạch trình bày", "30 giây",
    "Giúp Hội đồng biết cấu trúc và thời điểm xuất hiện của phần demo.",
    [
        "Bài trình bày gồm năm phần. Đầu tiên là bài toán và mục tiêu. Tiếp theo là giải pháp, kiến trúc và luồng nghiệp vụ. Phần thứ ba đi sâu vào gợi ý sản phẩm và phối cảnh AI. Sau đó nhóm trình bày kết quả kiểm chứng và demo. Cuối cùng là hạn chế và hướng phát triển.",
        "Nhóm đặt phần demo sau phần kiến trúc và kết quả kiểm thử để Hội đồng có đủ ngữ cảnh đánh giá những thao tác trên hệ thống.",
    ],
    "Từ cấu trúc trên, em xin bắt đầu với vấn đề thực tế mà đề tài muốn giải quyết.",
)

add_slide_section(
    doc, 3, "Bài toán mua nội thất trực tuyến", "1 phút 30 giây",
    "Chứng minh bài toán khó hơn một website thương mại điện tử thông thường.",
    [
        "Khi mua nội thất trực tuyến, khách hàng không chỉ quan tâm giá hay thông số. Họ cần hình dung một chiếc bàn, ghế hoặc tủ có phù hợp với căn phòng thật hay không. Ảnh sản phẩm riêng lẻ thường chưa trả lời được câu hỏi về màu sắc, tỷ lệ và sự hài hòa với không gian.",
        "Ở phía cửa hàng, dữ liệu sản phẩm, biến thể, tồn kho, đơn hàng, thanh toán và giao nhận phải được đồng bộ. Nếu các bước này tách rời, hệ thống dễ phát sinh tình trạng bán vượt tồn kho, cập nhật thanh toán không nhất quán hoặc mất dấu trạng thái giao hàng.",
        "Đối với nội thất, giao nhận cũng phức tạp hơn hàng hóa nhỏ. Đơn có thể cần phân công shipper, bằng chứng giao hàng, thu hộ, báo cáo sự cố và hẹn giao lại. Vì vậy, bài toán của đề tài gồm hai vế gắn chặt với nhau: hỗ trợ khách ra quyết định và bảo đảm quy trình vận hành xuyên suốt.",
        "Câu hỏi mà nhóm đặt ra là: có thể tích hợp AI vào đúng thời điểm khách hàng cần hình dung sản phẩm, nhưng vẫn giữ được một hệ thống bán lẻ hoàn chỉnh ở phía sau hay không?",
    ],
    "Từ bài toán đó, nhóm xác định mục tiêu và giới hạn triển khai như sau.",
)

add_slide_section(
    doc, 4, "Mục tiêu và phạm vi", "1 phút 15 giây",
    "Định vị đúng mức độ đóng góp và tránh tuyên bố vượt quá kết quả thực tế.",
    [
        "Nhóm đặt ra ba mục tiêu chính. Thứ nhất, hệ thống phục vụ ba vai trò là khách hàng, quản trị viên và nhân viên giao hàng. Thứ hai, dữ liệu phải đi xuyên suốt từ tìm kiếm sản phẩm đến khi giao hàng hoàn tất. Thứ ba, phối cảnh có hai chế độ: người dùng tự điều chỉnh và AI hỗ trợ tạo ảnh chân thực hơn.",
        "Trong phạm vi khóa luận, nhóm triển khai các nghiệp vụ thương mại điện tử chính, cơ chế gợi ý sản phẩm và phối cảnh trên ảnh phòng thật. Tiêu chí đánh giá là chức năng hoạt động đúng, các thành phần tích hợp được với nhau và quy trình AI có tính khả thi.",
        "Ngoài phạm vi của đề tài là huấn luyện một mô hình AI chuyên biệt hoàn toàn từ đầu, bảo đảm kích thước tuyệt đối trong mọi ảnh phòng, hoặc tối ưu hạ tầng cho quy mô người dùng rất lớn. Những nội dung này được xem là hướng phát triển, không phải kết quả đã hoàn thành.",
    ],
    "Sau khi xác định phạm vi, nhóm tổ chức hệ thống quanh ba vai trò cùng sử dụng một nguồn trạng thái đơn hàng.",
)

add_slide_section(
    doc, 5, "Ba vai trò trên cùng nền tảng", "1 phút 20 giây",
    "Làm rõ giá trị của trạng thái dùng chung thay vì liệt kê màn hình.",
    [
        "Vai trò thứ nhất là khách hàng. Khách có thể tìm kiếm, chọn biến thể, xem đánh giá, nhận gợi ý, thử sản phẩm trong phòng, đặt hàng, thanh toán và theo dõi sau mua.",
        "Vai trò thứ hai là quản trị viên. Admin quản lý sản phẩm, tồn kho, voucher, đơn hàng, doanh thu, hỗ trợ khách hàng, cấu hình gợi ý và điều phối giao hàng.",
        "Vai trò thứ ba là nhân viên giao hàng. Shipper quản lý ca làm, nhận đơn được phân công, cập nhật kết quả, tải bằng chứng giao hàng hoặc báo cáo sự cố và hẹn lại.",
        "Điểm quan trọng không phải là ba nhóm chức năng tách biệt. Cùng một đơn hàng được nhìn từ ba ngữ cảnh khác nhau nhưng sử dụng chung trạng thái nghiệp vụ. Khi một vai trò cập nhật đúng quy trình, hai vai trò còn lại có thể tiếp tục từ dữ liệu đó.",
    ],
    "Để thấy rõ sự liên kết này, em xin đi theo hành trình của một khách hàng cụ thể.",
    page_break=True,
)

add_slide_section(
    doc, 6, "Hành trình khách hàng", "1 phút 20 giây",
    "Diễn giải toàn bộ hệ thống bằng một ví dụ mua hàng dễ hình dung.",
    [
        "Giả sử khách hàng muốn mua một bộ bàn ăn. Đầu tiên, khách tìm kiếm bằng gợi ý từ khóa và lọc theo danh mục, màu, kích thước hoặc khoảng giá.",
        "Tại trang chi tiết, khách chọn đúng biến thể, kiểm tra tồn kho, xem đánh giá và các sản phẩm được gợi ý. Khi còn phân vân, khách có thể đưa sản phẩm vào ảnh phòng để xem thử, hoặc hỏi trợ lý AI về đặc điểm và cách phối hợp.",
        "Khi quyết định mua, khách thêm vào giỏ, áp dụng voucher, chọn địa chỉ được Mapbox chuẩn hóa và thanh toán bằng VNPay hoặc COD. Sau giao dịch, tồn kho, lịch sử đơn và trạng thái thanh toán được cập nhật.",
        "Cuối cùng, khách theo dõi quá trình giao nhận, nhận thông báo và có thể đánh giá hoặc yêu cầu hỗ trợ sau mua. Như vậy, hành trình không kết thúc ở nút đặt hàng mà kéo dài đến khi giao nhận và hậu mãi hoàn tất.",
    ],
    "Để thực hiện hành trình này, hệ thống được tách thành các lớp có trách nhiệm rõ ràng.",
)

add_slide_section(
    doc, 7, "Kiến trúc tổng thể", "1 phút 45 giây",
    "Giải thích luồng request từ client đến dữ liệu và dịch vụ AI.",
    [
        "Nhìn từ trái sang phải, Flutter là lớp giao diện cho khách hàng, Admin và Shipper. Client quản lý trạng thái hiển thị, gọi HTTP API và nhận các sự kiện thời gian thực qua SignalR.",
        "ASP.NET Core là trung tâm nghiệp vụ. Backend thực hiện xác thực JWT, phân quyền, kiểm tra dữ liệu, xử lý transaction và điều phối các dịch vụ bên ngoài. Entity Framework làm việc với SQL Server để lưu dữ liệu quan hệ như sản phẩm, đơn hàng, thanh toán và giao nhận.",
        "Các dịch vụ ngoài gồm VNPay, Mapbox, Gmail SMTP và nhà cung cấp AI. Với tác vụ xử lý ảnh, backend điều phối sang dịch vụ FastAPI sử dụng ONNX Runtime. Client không truy cập trực tiếp khóa bí mật của các dịch vụ này.",
        "Việc tách dịch vụ AI giúp phần suy luận có thời gian xử lý và nhu cầu tài nguyên riêng, không làm phức tạp toàn bộ backend. Đồng thời, ASP.NET Core vẫn kiểm soát quyền truy cập, dữ liệu đầu vào và URL kết quả trước khi trả về client.",
    ],
    "Mỗi công nghệ trong kiến trúc được chọn theo loại tải công việc mà nó đảm nhiệm.",
)

add_slide_section(
    doc, 8, "Lựa chọn công nghệ", "1 phút 10 giây",
    "Liên hệ công nghệ với trách nhiệm thay vì đọc một danh sách thư viện.",
    [
        "Flutter được chọn để xây dựng giao diện responsive từ một codebase và hỗ trợ trải nghiệm tương tác cao. ASP.NET Core phù hợp cho API phân quyền, transaction nghiệp vụ và tích hợp nhiều dịch vụ bên thứ ba.",
        "SQL Server lưu dữ liệu có quan hệ chặt chẽ, đặc biệt là sản phẩm, biến thể, tồn kho, đơn hàng và thanh toán. FastAPI cùng ONNX Runtime tách riêng các tác vụ xóa nền, phân vùng và ước lượng chiều sâu có thời gian xử lý dài hơn request thông thường.",
        "Ngoài ra, SignalR được dùng cho hỗ trợ và giao hàng theo thời gian thực. Các API như VNPay, Mapbox, OpenAI hoặc Gemini đều được gọi qua backend để bảo vệ khóa và kiểm soát luồng nghiệp vụ.",
    ],
    "Sau đây em xin đi sâu vào phần quan trọng nhất của backend: cách bảo toàn vòng đời đơn hàng.",
)

add_slide_section(
    doc, 9, "Vòng đời đơn hàng", "1 phút 20 giây",
    "Cho thấy transaction và trạng thái nghiệp vụ ngăn dữ liệu cập nhật nửa chừng.",
    [
        "Ở bước tạo đơn, backend kiểm tra tồn kho, tính voucher và phí giao hàng. Các thay đổi quan trọng được đặt trong transaction. Nếu một bước thất bại, hệ thống hoàn tác thay vì để tồn kho đã giảm nhưng đơn hàng chưa được tạo đầy đủ.",
        "Ở bước thanh toán, kết quả VNPay không được tin chỉ dựa trên tham số client gửi về. Backend kiểm tra chữ ký, mã giao dịch và trạng thái hiện tại của đơn trước khi ghi nhận thanh toán.",
        "Ở bước giao nhận, hệ thống đóng vòng trạng thái bằng việc phân công shipper, theo dõi bằng chứng giao hàng, thu hộ, sự cố và hẹn giao lại. Vì vậy, trạng thái 'đã giao' không phải một nút cập nhật đơn lẻ mà là kết quả của một chuỗi kiểm tra nghiệp vụ.",
    ],
    "Bên cạnh nghiệp vụ đơn hàng, hệ thống còn hỗ trợ khách ra quyết định bằng hai cơ chế gợi ý bổ sung cho nhau.",
)

add_slide_section(
    doc, 10, "Hai cơ chế gợi ý sản phẩm", "1 phút 30 giây",
    "Giải thích rõ vai trò của Apriori và content-based trong điều kiện dữ liệu khác nhau.",
    [
        "Cơ chế thứ nhất là Apriori. Thuật toán tìm các sản phẩm thường xuất hiện cùng nhau trong những đơn hợp lệ hoặc đã thanh toán. Support thể hiện mức độ phổ biến của một cặp sản phẩm; confidence thể hiện xác suất khách mua sản phẩm B khi đã mua sản phẩm A.",
        "Apriori phù hợp khi hệ thống đã có lịch sử giao dịch đủ lớn. Tuy nhiên, ở giai đoạn dữ liệu còn ít, số luật có thể hạn chế. Vì vậy nhóm bổ sung content-based để so khớp các đặc điểm như danh mục, màu sắc, kích thước và mức giá.",
        "Hai phương pháp giải quyết hai tình huống khác nhau: Apriori dựa trên hành vi mua chung, còn content-based dựa trên độ tương tự của sản phẩm. Admin có thể bật hoặc tắt từng phương pháp, thay đổi ngưỡng support và confidence, rồi xây lại luật khi dữ liệu tăng lên.",
        "Nhóm xem đây là nền tảng gợi ý có thể cấu hình, chưa khẳng định đã đạt mức cá nhân hóa cao vì chưa có đủ dữ liệu người dùng thực tế để đánh giá ngoại tuyến hoặc A/B testing.",
    ],
    "Nếu gợi ý giúp khách thu hẹp lựa chọn, thì phối cảnh AI giúp khách hình dung lựa chọn đó trong không gian thật.",
    presenter_note="Nếu chia hai người, đây là điểm chuyển người trình bày tự nhiên.",
)

add_slide_section(
    doc, 11, "Pipeline phối cảnh AI", "2 phút",
    "Trình bày đóng góp kỹ thuật nổi bật nhất theo đúng sáu bước.",
    [
        "Quy trình phối cảnh không chỉ chồng ảnh sản phẩm lên ảnh phòng. Bước một là xóa nền sản phẩm bằng U2NetP hoặc rembg để thu được vật thể và vùng alpha sạch.",
        "Bước hai là phân vùng ảnh phòng bằng SegFormer để nhận diện các khu vực như sàn và tường. Bước ba sử dụng Depth Anything V2 để ước lượng chiều sâu tương đối. Kết quả được tổng hợp thành lưới chiều sâu 32 x 24 nhằm hỗ trợ tỷ lệ gần - xa.",
        "Bước bốn là đặt vật thể lên mặt phẳng được chọn và điều chỉnh tỷ lệ theo phối cảnh. Bước năm hòa cảnh bằng cách ước lượng ánh sáng, bóng và màu sắc; hệ thống sử dụng lưới ánh sáng 16 x 12 làm dữ liệu tham khảo.",
        "Bước sáu xử lý che khuất. Vùng foreground và mask chỉnh sửa giúp một phần sản phẩm có thể nằm sau vật thể trong phòng thay vì luôn nổi lên trên ảnh.",
        "Sau sáu bước này, người dùng đã có một kết quả điều chỉnh được. OpenAI hoặc Gemini chỉ là lớp tạo ảnh tùy chọn để tăng mức chân thực, không phải điều kiện bắt buộc cho toàn bộ trình phối cảnh.",
    ],
    "Điểm quan trọng là AI hỗ trợ, nhưng người dùng vẫn giữ quyền kiểm soát kết quả.",
)

add_slide_section(
    doc, 12, "Người dùng kiểm soát phối cảnh", "1 phút 30 giây",
    "Giải thích hai chế độ và cơ chế sửa lỗi nhận diện.",
    [
        "Giao diện phối cảnh có hai chế độ bổ trợ. Ở chế độ tự điều chỉnh, người dùng thay đổi vị trí, kích thước, góc xoay, độ sáng và độ đậm bóng. Đây là chế độ có phản hồi trực tiếp và không phụ thuộc vào dịch vụ tạo sinh.",
        "Ở chế độ AI tạo ảnh, hệ thống gửi ảnh phòng và sản phẩm đã chuẩn bị qua backend để OpenAI hoặc Gemini hòa sản phẩm vào không gian. Kết quả này ưu tiên cảm giác chân thực nhưng có thể cần thời gian chờ và phụ thuộc quota dịch vụ ngoài.",
        "Khi nhận diện mặt phẳng chưa chính xác, người dùng có thể sửa bốn góc sàn hoặc tường. Khi che khuất sai, người dùng có thể điều chỉnh mask. Thiết kế sau đó được lưu, tải lại hoặc xuất ảnh.",
        "Vì vậy, nhóm định vị AI là phương án tham khảo. Hệ thống không che giấu sai số mà cung cấp công cụ để người dùng hiệu chỉnh trước khi sử dụng kết quả.",
    ],
    "Để các chức năng AI và thương mại điện tử hoạt động an toàn, bảo mật được đặt xuyên suốt luồng xử lý.",
)

add_slide_section(
    doc, 13, "Bảo mật và toàn vẹn", "1 phút 20 giây",
    "Nhóm biện pháp theo ba mục tiêu: xác thực, toàn vẹn giao dịch và an toàn dịch vụ AI.",
    [
        "Về xác thực, mật khẩu được băm bằng BCrypt; JWT có thời hạn và quyền được tách theo Customer, Admin và Shipper. Các endpoint quan trọng kiểm tra quyền thay vì chỉ ẩn nút trên giao diện.",
        "Về toàn vẹn giao dịch, tạo đơn sử dụng transaction, kết quả VNPay được xác minh chữ ký và các tài nguyên như thiết kế phòng được kiểm tra quyền sở hữu trước khi đọc hoặc sửa.",
        "Với thiết kế chưa đăng nhập, hệ thống dùng access token ngẫu nhiên có thời hạn 24 giờ. Ảnh kết quả sử dụng ETag và Cache-Control để giảm tải không cần thiết. Khóa OpenAI hoặc Gemini được giữ ở backend và nhà cung cấp có thể đổi bằng cấu hình.",
        "Như vậy, Flutter chỉ nhận dữ liệu và URL kết quả cần thiết; client không giữ khóa bí mật hoặc tự quyết định trạng thái nghiệp vụ quan trọng.",
    ],
    "Sau phần kỹ thuật, em xin quay lại các giao diện đã triển khai cho từng nhóm người dùng.",
)

add_slide_section(
    doc, 14, "Luồng chức năng khách hàng", "1 phút 10 giây",
    "Dùng giao diện thật để chứng minh luồng đã được triển khai, không chỉ dừng ở thiết kế.",
    [
        "Ở phía khách hàng, trang chi tiết hiển thị biến thể, giá, tồn kho, kích thước, đánh giá và sản phẩm gợi ý. Từ cùng màn hình, khách có thể thêm vào giỏ hoặc chuyển sang phối cảnh phòng.",
        "Trợ lý AI được đặt ngay trong trải nghiệm mua sắm và có thể đọc dữ liệu sản phẩm để trả lời về giá, đặc điểm, phối màu hoặc bố trí. Điểm khác với chatbot chung là câu hỏi được đặt trong đúng ngữ cảnh sản phẩm mà khách đang xem.",
        "Sau đó, khách tiếp tục qua giỏ hàng, voucher, địa chỉ, thanh toán, lịch sử đơn và hỗ trợ. Các ảnh trên slide là giao diện thật trong báo cáo và hệ thống triển khai.",
    ],
    "Khi khách đặt hàng xong, vòng vận hành được tiếp tục bởi Admin và Shipper.",
    page_break=True,
)

add_slide_section(
    doc, 15, "Admin và Shipper", "1 phút 20 giây",
    "Mô tả vòng trạng thái sau đặt hàng và cách hai vai trò phối hợp.",
    [
        "Admin theo dõi tổng quan, xác nhận đơn và thực hiện điều phối. Tại trung tâm điều phối, quản trị viên có thể xem tải giao hàng, phân ca, phân công shipper và tiếp nhận các sự cố.",
        "Shipper nhận danh sách đơn được giao, xem thông tin khách hàng, cập nhật trạng thái và tải bằng chứng. Nếu giao không thành công, shipper gửi lý do và đề xuất hẹn lại thay vì tự ý kết thúc đơn.",
        "Admin nhận thông tin sự cố để xử lý bước tiếp theo. SignalR cập nhật các sự kiện hỗ trợ và giao hàng theo thời gian thực, giúp giảm việc phải tải lại toàn bộ màn hình.",
        "Như vậy, hệ thống khép kín vòng vận hành từ xác nhận, phân công, giao hàng đến xử lý ngoại lệ.",
    ],
    "Nhóm không chỉ kiểm tra bằng giao diện mà còn chạy lại quá trình build và kiểm thử tự động.",
)

add_slide_section(
    doc, 16, "Kết quả triển khai và kiểm thử", "1 phút 20 giây",
    "Nêu số liệu có thể kiểm chứng và đồng thời nói rõ phạm vi của kiểm thử.",
    [
        "Từ source code trong workspace, hệ thống backend có 19 controller, 116 endpoint HTTP và 19 thực thể dữ liệu trong DbContext. Đây là số liệu được đếm trực tiếp, không phải ước lượng.",
        "Quá trình build backend thành công với 0 cảnh báo và 0 lỗi. Bộ kiểm thử tự động gồm 40 test Flutter và 4 test xử lý ảnh AI đều đạt tại thời điểm nhóm tạo bài trình bày.",
        "Các kết quả này chứng minh phạm vi chức năng và một số logic chính có thể chạy lặp lại. Tuy nhiên, nhóm không xem đây là bằng chứng cho khả năng chịu tải lớn hoặc an toàn bảo mật tuyệt đối, vì chưa thực hiện load test và kiểm thử xâm nhập chuyên sâu.",
        "Sau đây, nhóm xin demo ba điểm thể hiện rõ nhất giá trị của hệ thống.",
    ],
    "Em xin chuyển sang môi trường demo. Trong trường hợp dịch vụ ngoài mất kết nối, nhóm sẽ sử dụng các ảnh giao diện đã chuẩn bị trong slide.",
)

# Special demo section.
doc.add_page_break()
h = doc.add_paragraph(style="Heading 1")
r = h.add_run("SLIDE 17 | KỊCH BẢN DEMO HỆ THỐNG")
set_run_font(r, size=16, color=BLUE, bold=True)
add_label_paragraph(doc, "Thời lượng", "Khoảng 5 phút 30 giây")
add_callout(doc, "Mục tiêu", "Chứng minh ba luồng: khách hàng ra quyết định, đặt hàng tạo dữ liệu nghiệp vụ và vận hành tiếp tục ở Admin/Shipper.", fill=LIGHT_GRAY)

h2 = doc.add_paragraph(style="Heading 2")
r = h2.add_run("Chuẩn bị trước khi demo")
set_run_font(r, size=13, color=DARK, bold=True)
add_action(doc, "Mở sẵn ba tab hoặc ba cửa sổ: khách hàng, Admin và Shipper. Đăng nhập trước; chuẩn bị một sản phẩm bàn ăn có tồn kho, một voucher hợp lệ và một địa chỉ dễ tìm trên Mapbox.")
add_action(doc, "Giữ sẵn ảnh giao diện ở slide 12, 14 và 15. Chỉ mở sandbox VNPay hoặc gọi AI tạo sinh nếu kết nối ổn định.")
add_body(doc, "Lời mở demo: Trong phần demo, nhóm sẽ không đi qua tất cả màn hình mà tập trung vào ba điểm: hỗ trợ khách ra quyết định, bảo toàn dữ liệu khi đặt hàng và tiếp tục vòng vận hành sau đặt hàng.")

h3 = doc.add_paragraph(style="Heading 3")
r = h3.add_run("Phần 1 - Khách hàng chọn và xem thử sản phẩm")
set_run_font(r, size=12, color=RGBColor(31, 77, 120), bold=True)
add_action(doc, "Ở tab khách hàng, tìm sản phẩm bàn ăn; mở trang chi tiết; chọn màu hoặc kích thước; chỉ vào tồn kho, đánh giá và khu vực sản phẩm gợi ý.")
add_body(doc, "Tại đây, khách không chỉ xem ảnh và giá mà còn chọn đúng biến thể, kiểm tra tồn kho, đọc đánh giá và nhận gợi ý. Các dữ liệu này được lấy từ backend nên cùng một biến thể sẽ được dùng xuyên suốt khi thêm vào giỏ và tạo đơn.")
add_action(doc, "Chọn Room Design hoặc chức năng xem thử trong phòng; tải hoặc chọn ảnh phòng; đặt sản phẩm lên sàn; kéo một thanh kích thước hoặc xoay sản phẩm.")
add_body(doc, "Đây là chế độ tự điều chỉnh. Người dùng có thể thay đổi vị trí, tỷ lệ, góc xoay, độ sáng và bóng. Nếu mặt phẳng nhận diện chưa đúng, người dùng vẫn có thể sửa trước khi lưu hoặc tạo ảnh bằng AI.")

h3 = doc.add_paragraph(style="Heading 3")
r = h3.add_run("Phần 2 - Giỏ hàng, địa chỉ và thanh toán")
set_run_font(r, size=12, color=RGBColor(31, 77, 120), bold=True)
add_action(doc, "Thêm sản phẩm vào giỏ; thay đổi số lượng nếu phù hợp; áp dụng voucher; chuyển sang bước địa chỉ.")
add_body(doc, "Khi thêm vào giỏ, hệ thống giữ đúng biến thể đã chọn. Voucher và phí giao hàng được backend kiểm tra lại, thay vì chỉ tin số tiền hiển thị trên client.")
add_action(doc, "Nhập một phần địa chỉ để hiển thị gợi ý Mapbox; chọn địa chỉ chuẩn hóa. Chọn COD để bảo đảm demo nhanh, hoặc mở VNPay sandbox nếu đã chuẩn bị ổn định.")
add_body(doc, "Địa chỉ được chuẩn hóa để giảm sai lệch khi giao hàng. Khi tạo đơn, backend kiểm tra tồn kho và thực hiện transaction. Nếu thanh toán VNPay, kết quả còn phải được xác minh chữ ký trước khi ghi nhận.")
add_action(doc, "Hoàn tất đặt hàng và mở lịch sử đơn; chỉ vào mã đơn và trạng thái mới.")

h3 = doc.add_paragraph(style="Heading 3")
r = h3.add_run("Phần 3 - Admin phân công và Shipper xử lý")
set_run_font(r, size=12, color=RGBColor(31, 77, 120), bold=True)
add_action(doc, "Chuyển sang tab Admin; tìm đơn vừa tạo; xác nhận và mở trung tâm điều phối; phân công cho shipper đã chuẩn bị.")
add_body(doc, "Sau khi khách đặt hàng, Admin tiếp nhận cùng một đơn và thực hiện bước xác nhận, điều phối. Đây là điểm dữ liệu từ trải nghiệm mua sắm chuyển sang vận hành.")
add_action(doc, "Chuyển sang tab Shipper; mở đơn được phân công; cập nhật một trạng thái hoặc mở chức năng báo cáo sự cố.")
add_body(doc, "Shipper có thể nhận đơn, cập nhật kết quả, tải bằng chứng hoặc báo sự cố. Nếu có sự cố, Admin sẽ tiếp nhận để hẹn lại; đơn không bị kết thúc bằng một trạng thái thiếu thông tin.")
add_callout(doc, "Kết thúc demo", "Qua ba phần vừa rồi, nhóm đã chứng minh AI được đặt trong hành trình ra quyết định, còn dữ liệu đơn hàng vẫn tiếp tục xuyên suốt đến khâu vận hành.")

add_slide_section(
    doc, 18, "Hạn chế và hướng phát triển", "1 phút 20 giây",
    "Trình bày thẳng giới hạn, mỗi giới hạn gắn với một bước tiếp theo cụ thể.",
    [
        "Hạn chế thứ nhất là độ chính xác phối cảnh còn phụ thuộc góc chụp, ánh sáng và chất lượng ảnh. Hệ thống chưa bảo đảm tỷ lệ kích thước tuyệt đối. Hướng phát triển là bổ sung hiệu chuẩn camera, dữ liệu kiểm thử đa dạng hơn và nghiên cứu 3D hoặc AR.",
        "Hạn chế thứ hai là cá nhân hóa gợi ý. Apriori cần thêm lịch sử giao dịch, còn content-based chưa phản ánh đầy đủ sở thích từng người. Bước tiếp theo là thu thập lượt xem, tìm kiếm và giỏ hàng, sau đó thử lọc cộng tác hoặc mô hình lai.",
        "Hạn chế thứ ba là mức độ sẵn sàng vận hành lớn. Hệ thống còn phụ thuộc quota dịch vụ ngoài và chưa kiểm thử tải, bảo mật chuyên sâu. Hướng phát triển là triển khai cloud, bổ sung giám sát, retry, quản lý secret và tối ưu định tuyến giao hàng.",
        "Nhóm trình bày các hạn chế này để định vị đúng kết quả hiện tại và xác định những tiêu chí có thể kiểm chứng ở giai đoạn tiếp theo.",
    ],
    "Từ các kết quả và giới hạn trên, nhóm rút ra ba giá trị chính của đề tài.",
    page_break=True,
)

add_slide_section(
    doc, 19, "Kết luận", "50 giây",
    "Trả lời trực tiếp bài toán mở đầu bằng ba kết luận: người dùng, kỹ thuật và vận hành.",
    [
        "Về người dùng, hệ thống kết nối tìm kiếm, lựa chọn, xem thử và đặt hàng trong một hành trình. AI xuất hiện ở đúng thời điểm khách cần giảm sự không chắc chắn khi chọn nội thất.",
        "Về kỹ thuật, đề tài xây dựng kiến trúc nhiều dịch vụ, có phân quyền, transaction và pipeline xử lý ảnh có thể tiếp tục mở rộng.",
        "Về vận hành, Admin và Shipper cùng quản lý trạng thái đến khi giao hàng hoàn tất, bao gồm cả trường hợp ngoại lệ.",
        "Như vậy, My Furniture Store không tách AI thành một tính năng trình diễn. Hệ thống đặt AI vào quy trình bán lẻ hoàn chỉnh và tạo nền tảng khả thi để tiếp tục đánh giá bằng dữ liệu người dùng thực tế.",
    ],
    "Phần trình bày của nhóm đến đây là kết thúc. Nhóm xin cảm ơn Hội đồng.",
)

add_slide_section(
    doc, 20, "Cảm ơn và mời câu hỏi", "Phần thời gian còn lại",
    "Kết thúc ngắn gọn, giữ tư thế sẵn sàng trả lời.",
    [
        "Nhóm em xin trân trọng cảm ơn quý Thầy, Cô trong Hội đồng đã lắng nghe. Nhóm xin tiếp nhận câu hỏi và góp ý của Hội đồng.",
        "Khi nhận câu hỏi, em sẽ nhắc lại ý chính để xác nhận mình hiểu đúng, sau đó trả lời bằng quyết định thiết kế, bằng chứng trong source code hoặc demo, và giới hạn hiện tại nếu có.",
    ],
    None,
)


# Q&A appendix.
doc.add_page_break()
h = doc.add_paragraph(style="Heading 1")
r = h.add_run("PHỤ LỤC | CÂU HỎI HỘI ĐỒNG CÓ THỂ ĐẶT")
set_run_font(r, size=16, color=BLUE, bold=True)
add_callout(doc, "Cách dùng", "Không học thuộc từng chữ. Hãy nhớ cấu trúc của câu trả lời và chỉ nêu những gì nhóm có bằng chứng.", fill=LIGHT_GRAY)

qa_items = [
    (
        "1. Điểm mới hoặc đóng góp chính của đề tài là gì?",
        "Đóng góp chính không nằm ở việc tạo ra một mô hình AI mới. Điểm của đề tài là tích hợp phối cảnh AI và gợi ý sản phẩm vào một quy trình bán lẻ nội thất hoàn chỉnh, đồng thời giữ được transaction đơn hàng, thanh toán, giao nhận và phân quyền. Nhóm cũng triển khai cơ chế người dùng chỉnh sửa kết quả AI thay vì coi kết quả nhận diện là tuyệt đối.",
    ),
    (
        "2. Vì sao chọn Apriori thay vì một mô hình học sâu?",
        "Apriori phù hợp để tạo luật mua kèm có thể giải thích và dễ cấu hình khi dữ liệu giao dịch có cấu trúc. Với giai đoạn hiện tại, dữ liệu người dùng chưa đủ để huấn luyện và đánh giá một mô hình học sâu đáng tin cậy. Nhóm bổ sung content-based để giảm vấn đề cold start, và xem lọc cộng tác hoặc mô hình lai là hướng phát triển khi có thêm dữ liệu.",
    ),
    (
        "3. Độ chính xác của phối cảnh được đánh giá như thế nào?",
        "Nhóm đánh giá tính khả thi của quy trình: xóa nền, nhận diện sàn hoặc tường, ước lượng chiều sâu, đặt vật thể, hòa sáng và xử lý che khuất. Hệ thống chưa có bộ dữ liệu chuẩn để đo sai số kích thước tuyệt đối. Vì vậy, nhóm không tuyên bố độ chính xác tuyệt đối và cung cấp cơ chế người dùng sửa mặt phẳng, mask và thông số hiển thị.",
    ),
    (
        "4. Vì sao tách FastAPI khỏi ASP.NET Core?",
        "FastAPI phù hợp cho các thư viện xử lý ảnh và ONNX trong hệ sinh thái Python. Tách riêng giúp tài nguyên suy luận, thời gian chờ và phụ thuộc mô hình không làm phức tạp backend nghiệp vụ. ASP.NET Core vẫn là điểm kiểm soát xác thực, quyền truy cập, dữ liệu đầu vào và URL kết quả.",
    ),
    (
        "5. Nếu OpenAI hoặc Gemini không hoạt động thì hệ thống có dùng được không?",
        "Có. Chế độ tự điều chỉnh vẫn hoạt động với kết quả phân tích ảnh, vị trí, tỷ lệ, xoay, sáng, bóng và mask. OpenAI hoặc Gemini là lớp tạo ảnh tùy chọn để tăng mức chân thực. Hệ thống có thể đổi nhà cung cấp bằng cấu hình và nên bổ sung retry, quota monitoring ở giai đoạn vận hành lớn.",
    ),
    (
        "6. Hệ thống bảo vệ khóa API và dữ liệu người dùng như thế nào?",
        "Khóa dịch vụ được giữ ở backend; Flutter không truy cập trực tiếp. Hệ thống dùng JWT có thời hạn, BCrypt cho mật khẩu, phân quyền theo vai trò, kiểm tra quyền sở hữu thiết kế, xác minh chữ ký VNPay và access token ngẫu nhiên có thời hạn cho thiết kế chưa đăng nhập.",
    ),
    (
        "7. Các con số 19 controller, 116 endpoint và 44 test có ý nghĩa gì?",
        "Các số liệu này chứng minh phạm vi triển khai và mức độ có thể kiểm chứng của source code. Build backend thành công và 44 test hiện có đều đạt. Tuy nhiên, số lượng test không đồng nghĩa hệ thống đã được kiểm thử tải hoặc kiểm thử bảo mật chuyên sâu; nhóm nêu rõ hai nội dung này chưa nằm trong kết quả hiện tại.",
    ),
    (
        "8. Điểm yếu lớn nhất của hệ thống hiện tại là gì?",
        "Điểm yếu lớn nhất là chưa có dữ liệu người dùng và bộ dữ liệu ảnh chuẩn đủ lớn để đánh giá định lượng chất lượng gợi ý và sai số phối cảnh. Do đó, kết quả hiện tại là nền tảng kỹ thuật khả thi; bước tiếp theo cần tập trung vào dữ liệu thực tế, chỉ số đánh giá và thử nghiệm người dùng.",
    ),
    (
        "9. Nếu hai khách cùng mua sản phẩm cuối cùng thì xử lý thế nào?",
        "Backend kiểm tra tồn kho khi tạo đơn và sử dụng transaction để tránh cập nhật nửa chừng. Trong triển khai thực tế quy mô lớn, nhóm sẽ bổ sung chiến lược khóa hoặc concurrency token phù hợp, cơ chế giữ hàng theo thời hạn thanh toán và kiểm thử cạnh tranh ở mức cơ sở dữ liệu.",
    ),
    (
        "10. Hướng phát triển ưu tiên nhất là gì?",
        "Ưu tiên thứ nhất là xây bộ dữ liệu đánh giá phối cảnh và thu thập hành vi người dùng để có chỉ số định lượng. Sau đó là cải thiện cá nhân hóa bằng mô hình lai, bổ sung cloud monitoring, retry, quản lý secret và kiểm thử tải. 3D hoặc AR là hướng dài hạn sau khi nền tảng dữ liệu và vận hành ổn định.",
    ),
]

for question, answer in qa_items:
    h2 = doc.add_paragraph(style="Heading 2")
    r = h2.add_run(question)
    set_run_font(r, size=13, color=DARK, bold=True)
    add_body(doc, answer)

h = doc.add_paragraph(style="Heading 1")
r = h.add_run("CHECKLIST TRƯỚC BUỔI BẢO VỆ")
set_run_font(r, size=16, color=BLUE, bold=True)
check_items = [
    ("Tệp trình chiếu", "Mở đúng bản có trang đầu tối giản; kiểm tra font và tỷ lệ 16:9 trên máy trình chiếu."),
    ("Môi trường demo", "Chạy backend, frontend, SQL Server và dịch vụ AI; kiểm tra CORS, URL API và tài khoản ba vai trò."),
    ("Dữ liệu demo", "Chuẩn bị sản phẩm còn tồn kho, voucher hợp lệ, shipper đang hoạt động và một ảnh phòng có sàn rõ."),
    ("Dịch vụ ngoài", "Kiểm tra Mapbox, SignalR và sandbox VNPay; nếu không ổn định thì dùng COD và ảnh dự phòng."),
    ("Thời gian", "Chạy thử toàn bộ bài với đồng hồ; nếu vượt thời gian, rút gọn slide 5, 8, 13 và 14 trước, không cắt phần demo hoặc kết luận."),
    ("Phân công", "Thống nhất ai nói phần nào, ai thao tác demo, ai tiếp nhận câu hỏi kỹ thuật và ai tiếp nhận câu hỏi nghiệp vụ."),
]
for label, value in check_items:
    add_label_paragraph(doc, label, value, after=6)

h = doc.add_paragraph(style="Heading 1")
r = h.add_run("NGUỒN XÂY DỰNG KỊCH BẢN")
set_run_font(r, size=16, color=BLUE, bold=True)
add_body(doc, "Báo cáo khóa luận: BaoCaoKhoaLuan_Final1.docx.")
add_body(doc, "Bộ slide bảo vệ: BaoCao_BaoVe_KhoaLuan_MyFurnitureStore_TrangDau_ToiGian.pptx.")
add_body(doc, "Source code backend tại D:\\KLTN\\furniture_BE và frontend tại D:\\KLTN\\khoaluan.")

doc.core_properties.title = "Kịch bản thuyết trình khóa luận - My Furniture Store"
doc.core_properties.subject = "Kịch bản trình bày khoảng 30 phút cho buổi bảo vệ khóa luận"
doc.core_properties.author = "Đỗ Minh Vĩ Khang và Trần Ngọc Nhật Duy"
doc.core_properties.keywords = "khóa luận, kịch bản thuyết trình, nội thất, AI, My Furniture Store"

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
